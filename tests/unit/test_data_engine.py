"""WS43-F26 — the data engine of the coding sandbox (WS-43y1a).

Spec: ``project-docs/specs/maf_coding_engine.md`` §7.10 ("Reading a messy
file", "Types", "CSV files", "The SQL rule", "The result envelope"), the
WS-43y1 slice and §10 (the fence row WS43-F26).

The engine is ``apps/services/orchestrator/sandbox/data_engine.py``. It runs in
the sandbox image, so these tests load it by its path, as the image does. They
run it on the host Python with the host's DuckDB, at the pin of the image lock.

The fixtures sit in this file, beside their test (``tests/fixtures/README.md``
rule 1). A text fixture needs no script to make it, and the tests make each
encoded copy (UTF-16, cp1252) from the text.

**The Docker half** carries the ``sandbox_docker`` marker. It runs each verb in
the image as uid 1000, read-only and with no network. It also runs each lock
case there, with the statement check off.
"""

from __future__ import annotations

import ast
import csv
import importlib.util
import io
import json
import os
import re
import stat
import sys
import tempfile
import time
import tracemalloc
import uuid
import zipfile
from decimal import Decimal
from pathlib import Path
from types import ModuleType
from typing import Any

import duckdb
import pytest

from tests.unit import _xlsx_build as xb
from tests.unit._xlsx_build import inline, num
from tests.unit._xlsx_build import row as xrow
from tests.unit.test_coding_sandbox_image import (  # noqa: F401 — a fixture, used by name
    _DOCKERFILE,
    _RUN_FLAGS,
    _docker,
    _instructions,
    coding_sandbox_image,
)

_REPO = Path(__file__).resolve().parents[2]
_ENGINE_PATH = _REPO / "apps" / "services" / "orchestrator" / "sandbox" / "data_engine.py"
_IMAGE_LOCK = _ENGINE_PATH.parent / "requirements.txt"


def _load_engine() -> ModuleType:
    spec = importlib.util.spec_from_file_location("data_engine_under_test", _ENGINE_PATH)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module  # dataclasses read the module of a class
    spec.loader.exec_module(module)
    return module


E = _load_engine()

# --------------------------------------------------------------------------
# The fixtures.
# --------------------------------------------------------------------------

INVOICES = """\
Invoice,Customer,Amount,Tax %,Due date,Paid,Ref

INV-001,Acme,"₹1,23,456.50",18%,01/03/2026,yes,007
INV-002,Beta,"$1,234.00",5%,15/03/2026,no,008
INV-003,Gamma,(500),0%,31/03/2026,yes,009
INV-004,Delta,250.25,12.5%,02/04/2026,no,010
"""
AMOUNT_SUM = "124440.75"

EU_SALES = "Region;Revenue;Units\nNorth;1.234,5;10\nSouth;987,25;3\nEast;12,5;7\n"


@pytest.fixture
def dirs(tmp_path: Path) -> Any:
    return E.Dirs(run=tmp_path / "run", data=tmp_path / "data", outputs=tmp_path / "outputs")


def _call(dirs: Any, verb: str, **req: Any) -> dict[str, Any]:
    """Write a request file, run the engine's main, and read its answer file."""
    rid = uuid.uuid4().hex
    (dirs.run / "requests").mkdir(parents=True, exist_ok=True)
    (dirs.run / "requests" / f"{rid}.json").write_text(json.dumps(req), encoding="utf-8")
    assert E.main([verb, rid], dirs) == 0
    return json.loads((dirs.run / "answers" / f"{rid}.json").read_text(encoding="utf-8"))


def _load(dirs: Any, name: str, content: str | bytes, **opts: Any) -> dict[str, Any]:
    (dirs.run / "src").mkdir(parents=True, exist_ok=True)
    data = content.encode("utf-8") if isinstance(content, str) else content
    (dirs.run / "src" / name).write_bytes(data)
    return _call(dirs, "load", source=f"src/{name}", **opts)


def _ok(answer: dict[str, Any]) -> dict[str, Any]:
    assert answer["ok"], answer
    return answer


def _columns(dirs: Any, ds_id: str) -> dict[str, dict[str, Any]]:
    manifest = json.loads((dirs.data / ds_id / "manifest.json").read_text(encoding="utf-8"))
    return {c["name"]: c for c in manifest["tables"][0]["columns"]}


@pytest.fixture
def invoices(dirs: Any) -> str:
    return _ok(_load(dirs, "Invoices.csv", INVOICES))["dataset_id"]


# --------------------------------------------------------------------------
# Done-when 1: the header, the left-out rows, the findings and the types.
# --------------------------------------------------------------------------


def test_a_csv_gets_its_header_its_types_and_its_findings(dirs: Any) -> None:
    answer = _ok(_load(dirs, "Invoices.csv", INVOICES))
    assert re.fullmatch(r"ds_[0-9a-f]{32}", answer["dataset_id"])  # 128 bits
    [table] = answer["tables"]
    assert table["name"] == "invoices" and table["rows"] == 4
    assert table["range"] == "invoices!A3:G6"
    assert {"what": "header rows", "range": "invoices!A1:G1"} in table["left_out"]
    assert {"what": "empty rows", "count": 1, "rows": [2]} in table["left_out"]
    cols = _columns(dirs, answer["dataset_id"])
    types = {name: c["type"] for name, c in cols.items()}
    assert types == {
        "Invoice": "text", "Customer": "text", "Amount": "decimal", "Tax %": "decimal",
        "Due date": "date", "Paid": "boolean", "Ref": "text",
    }  # fmt: skip
    assert cols["Amount"]["parsed_from_text"] == 3  # ₹ Indian, $ western, (500)
    assert cols["Amount"]["min"] == "-500.00" and cols["Amount"]["max"] == "123456.50"
    assert cols["Tax %"]["findings"] == ["percent_values"]
    assert cols["Due date"]["min"] == "2026-03-01" and cols["Due date"]["findings"] == []
    assert cols["Ref"]["samples"] == ["007", "008", "009"]  # a zero first stays text


def test_the_sum_is_exact_and_names_its_source(dirs: Any, invoices: str) -> None:
    answer = _ok(
        _call(dirs, "query", dataset_id=invoices, sql='SELECT sum("Amount") AS total FROM invoices')
    )
    assert answer["rows"] == [[AMOUNT_SUM]]
    assert answer["rows_scanned"] == {"invoices": 4}
    [source] = answer["source_ranges"]
    assert source["table"] == "invoices" and source["range"] == "invoices!A3:G6"
    assert source["left_out"][0]["what"] == "header rows"
    for field in ("dataset_id", "table", "query", "rows_scanned", "rows_returned", "truncated",
                  "truncated_reason", "columns", "rows", "source_ranges", "elapsed_ms"):  # fmt: skip
        assert field in answer, field
    assert answer["query"].endswith("LIMIT 101") and not answer["truncated"]


def test_a_semicolon_file_reads_a_decimal_comma(dirs: Any) -> None:
    answer = _ok(_load(dirs, "eu.csv", EU_SALES))
    manifest = json.loads((dirs.data / answer["dataset_id"] / "manifest.json").read_text("utf-8"))
    assert manifest["delimiter"] == ";" and manifest["findings"] == ["decimal_comma"]
    total = _ok(
        _call(dirs, "query", dataset_id=answer["dataset_id"], sql="SELECT sum(Revenue) FROM eu")
    )
    assert total["rows"] == [["2234.25"]]
    assert _columns(dirs, answer["dataset_id"])["Units"]["type"] == "integer"


@pytest.mark.parametrize("encoding", ["utf-8-sig", "utf-16", "cp1252"])
def test_each_encoding_reads_the_same(dirs: Any, encoding: str) -> None:
    answer = _ok(_load(dirs, "e.csv", "Name,Amount\nCafé €,1\nTea,2\n".encode(encoding)))
    cols = _columns(dirs, answer["dataset_id"])
    assert cols["Name"]["samples"] == ["Café €", "Tea"] and cols["Amount"]["type"] == "integer"


@pytest.mark.parametrize(
    ("name", "text"), [("t.tsv", "a\tb\n1\t2\n"), ("p.csv", "a|b\n1|2\n3|4\n")]
)
def test_a_tab_and_a_pipe_delimiter(dirs: Any, name: str, text: str) -> None:
    answer = _ok(_load(dirs, name, text))
    assert [c["name"] for c in answer["tables"][0]["columns"]] == ["a", "b"]


def test_a_mixed_column_stays_text_and_counts_each_type(dirs: Any) -> None:
    answer = _ok(_load(dirs, "m.csv", "Code,Note\n1,a\nx,b\n2026-01-05,c\nyes,d\n"))
    code = _columns(dirs, answer["dataset_id"])["Code"]
    assert code["type"] == "text"
    assert code["type_counts"] == {"number": 1, "boolean": 1, "date": 1, "text": 1}


def test_a_serial_date_column_gets_a_finding_and_no_conversion(dirs: Any) -> None:
    answer = _ok(_load(dirs, "s.csv", "Order date,Qty\n45992,3\n46000,4\n"))
    cols = _columns(dirs, answer["dataset_id"])
    assert cols["Order date"]["type"] == "integer"
    assert cols["Order date"]["findings"] == ["possible_excel_serial_date"]
    assert cols["Qty"]["findings"] == []


def test_dates_times_and_an_ambiguous_day_month(dirs: Any) -> None:
    text = "When,At\n01/02/2026,2026-03-01 10:00:00\n03/04/2026,2026-03-02 11:30:00\n"
    cols = _columns(dirs, _ok(_load(dirs, "d.csv", text))["dataset_id"])
    assert cols["When"]["type"] == "date" and cols["When"]["findings"] == ["ambiguous_day_month"]
    assert cols["When"]["min"] == "2026-02-01"  # day first
    assert cols["At"]["type"] == "datetime"


def test_two_header_rows_join_their_levels(dirs: Any) -> None:
    answer = _ok(_load(dirs, "h.csv", "Sales,Sales,Cost\nQ1,Q2,Q1\n1,2,3\n", header_rows=2))
    names = [c["name"] for c in answer["tables"][0]["columns"]]
    assert names == ["Sales / Q1", "Sales / Q2", "Cost / Q1"]


def test_a_second_load_reads_nothing_and_the_options_change_the_id(dirs: Any) -> None:
    first = _ok(_load(dirs, "Invoices.csv", INVOICES))
    again = _ok(_load(dirs, "Invoices.csv", INVOICES))
    assert again["reused"] and again["dataset_id"] == first["dataset_id"]
    other = _ok(_load(dirs, "Invoices.csv", INVOICES, header_rows=2))
    assert other["dataset_id"] != first["dataset_id"]
    assert not any(p.name.startswith(".stage-") for p in dirs.data.iterdir())


@pytest.mark.parametrize(
    ("name", "says"),
    [("a.xls", "old .xls"), ("a.xlsb", ".xlsb"), ("a.ods", ".ods"),
     ("a.txt", ".csv, .tsv, .xlsx and .xlsm")],
)  # fmt: skip
def test_the_engine_refuses_a_kind_it_does_not_read_with_a_reason(
    dirs: Any, name: str, says: str
) -> None:
    answer = _load(dirs, name, b"PK\x03\x04 not read")
    assert answer["error"] == "unsupported_kind" and says in answer["message"]


@pytest.mark.parametrize("source", ["../escape.csv", "src/../../escape.csv", "src/missing.csv", ""])
def test_a_source_outside_the_run_data_dir_is_refused(dirs: Any, source: str) -> None:
    (dirs.run.parent / "escape.csv").write_text("a\n1\n", encoding="utf-8")
    answer = _call(dirs, "load", source=source)
    assert not answer["ok"] and answer["error"] in ("bad_request", "not_found")


def test_the_size_and_row_caps_refuse_a_load(dirs: Any) -> None:
    assert _load(dirs, "a.csv", INVOICES, max_file_bytes=10)["error"] == "too_large"
    assert _load(dirs, "b.csv", INVOICES, max_rows=2)["error"] == "too_many_rows"
    assert _load(dirs, "c.csv", INVOICES, header_rows=4)["error"] == "bad_request"


# --------------------------------------------------------------------------
# Done-when 2 and 6: the statement check, and the lock with the check off.
# --------------------------------------------------------------------------


@pytest.mark.parametrize(
    "sql",
    [
        "SELECT 1; SELECT 2",
        "SELECT * FROM invoices; DROP VIEW invoices",
        "ATTACH 'other.db'",
        "COPY invoices TO 'w.csv'",
        "INSTALL httpfs",
        "LOAD httpfs",
        "SET threads = 4",
        "RESET enable_external_access",
        "PRAGMA version",
        "PRAGMA threads = 4",
        "CREATE TABLE z AS SELECT 1",
        "EXPORT DATABASE 'x'",
        "CALL pragma_version()",
        "DESCRIBE invoices",
    ],
)
def test_the_statement_check_refuses(dirs: Any, invoices: str, sql: str) -> None:
    answer = _call(dirs, "query", dataset_id=invoices, sql=sql)
    assert answer["error"] in ("one_select", "sql_error"), answer
    with pytest.raises(E.Refused):
        E._check_sql(sql)


def test_a_second_statement_is_refused_by_name(dirs: Any, invoices: str) -> None:
    """M3: the refusal is the statement check's, not a parse error of the wrap."""
    answer = _call(dirs, "query", dataset_id=invoices, sql="SELECT 1; SELECT 2")
    assert answer["error"] == "one_select"
    assert E._check_sql("WITH x AS (SELECT 1 AS a) SELECT a FROM x -- note") == (
        "WITH x AS (SELECT 1 AS a) SELECT a FROM x -- note"
    )


def _lock_cases(dirs: Any, invoices: str) -> dict[str, str]:
    other = (dirs.run / "other.csv").resolve().as_posix()
    (dirs.run / "other.csv").write_text("a\n1\n", encoding="utf-8")
    sibling = _ok(_load(dirs, "eu.csv", EU_SALES))["dataset_id"]
    ds = (dirs.data / invoices).resolve().as_posix()
    return {
        "read_csv": f"SELECT * FROM read_csv('{other}')",
        "read_text": f"SELECT * FROM read_text('{other}')",
        "glob": f"SELECT * FROM glob('{ds}/*')",
        "getenv": "SELECT getenv('PATH')",
        "sibling dataset": f"SELECT * FROM read_parquet('{(dirs.data / sibling).resolve().as_posix()}/eu.parquet')",
        "attach": f"ATTACH '{ds}/x.db'",
        "copy into the dataset dir": f"COPY (SELECT 1 AS a) TO '{ds}/w.csv'",
        "copy elsewhere": f"COPY (SELECT 1 AS a) TO '{dirs.run.resolve().as_posix()}/w.csv'",
        "install": "INSTALL httpfs",
        "load": "LOAD httpfs",
        "set": "SET threads = 4",
        "reset": "RESET enable_external_access",
        "pragma": "PRAGMA threads = 4",
        "export": f"EXPORT DATABASE '{ds}/exp'",
        "copy into the spill dir": f"COPY (SELECT 1 AS a) TO '{E._spill_dir().as_posix()}/w.csv'",
        **_overwrite_cases(f"{ds}/invoices.parquet"),
    }


def _overwrite_cases(parquet: str) -> dict[str, str]:
    """P1-A: every COPY form that could write the dataset's own Parquet file."""
    copy = f"COPY (SELECT 99 AS a) TO '{parquet}'"
    return {
        "overwrite parquet": f"{copy} (FORMAT parquet)",
        "overwrite parquet, no tmp file": f"{copy} (FORMAT parquet, USE_TMP_FILE false)",
        "overwrite parquet, tmp file": f"{copy} (FORMAT parquet, USE_TMP_FILE true)",
        "overwrite csv, no tmp file": f"{copy} (FORMAT csv, USE_TMP_FILE false)",
        "overwrite csv": f"{copy} (FORMAT csv)",
        "overwrite true": f"{copy} (FORMAT parquet, OVERWRITE true)",
        "overwrite or ignore": f"{copy} (FORMAT parquet, OVERWRITE_OR_IGNORE true)",
        "partition by": f"{copy} (FORMAT parquet, PARTITION_BY (a), OVERWRITE_OR_IGNORE)",
    }


def test_the_lock_stops_each_case_with_the_statement_check_off(dirs: Any, invoices: str) -> None:
    """Rule 4, done-when 2 and 6. The query runs on the locked connection itself."""
    before = {p.name: p.read_bytes() for p in (dirs.data / invoices).iterdir()}
    failures = []
    for label, sql in _lock_cases(dirs, invoices).items():
        con, _ = E._open_locked(dirs, invoices)
        try:
            con.execute(sql).fetchall()
            failures.append(label)
        except duckdb.Error:
            pass
        # The views still answer after the refusal, from the same bytes.
        total = con.execute('SELECT count(*), sum("Amount") FROM invoices').fetchone()
        assert total == (4, Decimal(AMOUNT_SUM)), label
    assert failures == []
    assert {p.name: p.read_bytes() for p in (dirs.data / invoices).iterdir()} == before
    assert not (dirs.run / "w.csv").exists()


def test_the_dataset_is_read_only_after_its_load(dirs: Any, invoices: str) -> None:
    """P1-A: no file of the dataset is writable, and on POSIX the dir is 0555."""
    for path in (dirs.data / invoices).iterdir():
        assert not os.access(path, os.W_OK), path
    if os.name == "posix":
        assert stat.S_IMODE((dirs.data / invoices).stat().st_mode) == 0o555


def test_the_locked_connection_holds_every_setting_of_rule_3(dirs: Any, invoices: str) -> None:
    """P2-G: read each setting back from the locked connection itself."""
    con, _ = E._open_locked(dirs, invoices)
    settings = dict(con.execute("SELECT name, value FROM duckdb_settings()").fetchall())
    parquet = (dirs.data / invoices / "invoices.parquet").resolve().as_posix()
    assert settings["allowed_paths"].strip("[]'") == parquet
    # DuckDB adds its spill dir to the allowed dirs. It is new for each process.
    assert settings["allowed_directories"].strip("[]'") == E._spill_dir().resolve().as_posix() + "/"
    for name, value in (
        ("autoinstall_known_extensions", "false"),
        ("autoload_known_extensions", "false"),
        ("enable_external_access", "false"),
        ("lock_configuration", "true"),
        ("threads", "1"),
    ):
        assert settings[name] == value, name


@pytest.mark.parametrize("sql", ["read_csv", "read_text", "glob", "getenv"])
def test_a_query_that_reads_a_file_is_refused_through_the_verb(
    dirs: Any, invoices: str, sql: str
) -> None:
    answer = _call(dirs, "query", dataset_id=invoices, sql=_lock_cases(dirs, invoices)[sql])
    assert answer["error"] == "sql_error", answer


def test_the_pinned_duckdb_has_allowed_paths() -> None:
    """Rule 5: with no ``allowed_paths``, the engine would need the fallback."""
    names = {
        row[0] for row in duckdb.connect().execute("SELECT name FROM duckdb_settings()").fetchall()
    }
    assert {"allowed_paths", "enable_external_access", "lock_configuration"} <= names


# --------------------------------------------------------------------------
# Done-when 3: the caps.
# --------------------------------------------------------------------------


def test_a_query_over_the_row_cap_sets_truncated(dirs: Any, invoices: str) -> None:
    answer = _ok(
        _call(dirs, "query", dataset_id=invoices, sql="SELECT * FROM invoices", max_result_rows=2)
    )
    assert answer["truncated"] and answer["truncated_reason"] == ["rows"]
    assert answer["rows_returned"] == 2 and len(answer["rows"]) == 2
    exact = _ok(
        _call(dirs, "query", dataset_id=invoices, sql="SELECT * FROM invoices", max_result_rows=4)
    )
    assert not exact["truncated"] and exact["rows_returned"] == 4


def test_a_long_cell_is_cut_and_says_so(dirs: Any, invoices: str) -> None:
    answer = _ok(_call(dirs, "query", dataset_id=invoices, sql="SELECT repeat('x', 300) AS s"))
    assert answer["truncated_reason"] == ["cells"] and len(answer["rows"][0][0]) == 200


def test_a_slow_query_stops_at_the_time_cap(dirs: Any, invoices: str) -> None:
    sql = "SELECT sum(a.range * b.range) FROM range(1000000) a, range(1000000) b"
    started = time.monotonic()
    answer = _call(dirs, "query", dataset_id=invoices, sql=sql, timeout_seconds=1)
    assert answer["error"] == "time", answer
    assert answer["truncated"] and answer["truncated_reason"] == ["time"] and answer["rows"] == []
    assert time.monotonic() - started < 15


def test_a_request_past_an_engine_cap_is_refused(dirs: Any, invoices: str) -> None:
    for field, value in (
        ("max_result_rows", 501),
        ("timeout_seconds", 0),
        ("max_result_rows", True),
    ):
        answer = _call(dirs, "query", dataset_id=invoices, sql="SELECT 1", **{field: value})
        assert answer["error"] == "bad_request", (field, value)


# --------------------------------------------------------------------------
# The other verbs.
# --------------------------------------------------------------------------


def test_a_row_with_its_source_row_gets_its_cells(dirs: Any, invoices: str) -> None:
    sql = 'SELECT _src_row, "Invoice" FROM invoices WHERE "Customer" = \'Beta\''
    answer = _ok(_call(dirs, "query", dataset_id=invoices, sql=sql))
    assert answer["rows"] == [[4, "INV-002"]] and answer["cells"] == ["invoices!A4:G4"]


def test_preview_gives_rows_in_source_order_with_cells(dirs: Any, invoices: str) -> None:
    answer = _ok(_call(dirs, "preview", dataset_id=invoices, table="invoices", offset=1, limit=2))
    assert [r[1] for r in answer["rows"]] == ["INV-002", "INV-003"]
    assert answer["cells"] == ["invoices!A4:G4", "invoices!A5:G5"]
    assert answer["truncated"] and answer["table"] == "invoices"
    assert (
        _call(dirs, "preview", dataset_id=invoices, table="invoices", limit=101)["error"]
        == "bad_request"
    )
    assert _call(dirs, "preview", dataset_id=invoices, table="nope")["error"] == "not_found"


def test_profile_gives_the_counts_and_the_number_statistics(dirs: Any, invoices: str) -> None:
    answer = _ok(_call(dirs, "profile", dataset_id=invoices, table="invoices", column="Amount"))
    row = dict(zip([c["name"] for c in answer["columns"]], answer["rows"][0], strict=True))
    assert row["rows"] == 4 and row["empty"] == 0 and row["distinct"] == 4
    assert row["minimum"] == "-500.00" and row["maximum"] == "123456.50"
    assert {"mean", "median", "p5_p95"} <= set(row)
    assert len(answer["top_values"]) == 4 and answer["rows_scanned"] == {"invoices": 4}
    assert (
        _call(dirs, "profile", dataset_id=invoices, table="invoices", column="x")["error"]
        == "not_found"
    )


def test_export_writes_every_row_and_answers_none(dirs: Any) -> None:
    ds = _ok(_load(dirs, "f.csv", "Name,Amount\n=1+1,2\nAda,3\n"))["dataset_id"]
    sql = "SELECT * EXCLUDE (_src_row) FROM f ORDER BY Amount"
    answer = _ok(_call(dirs, "export", dataset_id=ds, sql=sql, format="csv", name="out"))
    assert (
        answer["file"] == "outputs/out.csv"
        and answer["rows"] == []
        and answer["rows_returned"] == 2
    )
    text = (dirs.outputs / "out.csv").read_text(encoding="utf-8")
    assert text.splitlines() == ["Name,Amount", "'=1+1,2", "Ada,3"]
    again = _ok(_call(dirs, "export", dataset_id=ds, sql=sql, format="csv", name="out"))
    assert again["file"] == "outputs/out-2.csv"
    book = _ok(_call(dirs, "export", dataset_id=ds, sql=sql, format="xlsx", name="out"))
    from openpyxl import load_workbook

    sheet = load_workbook(dirs.outputs / "out.xlsx")["export"]
    assert sheet["A2"].value == "=1+1" and sheet["A2"].data_type == "s"
    assert book["rows_returned"] == 2
    bad = _call(dirs, "export", dataset_id=ds, sql="SELECT 1", format="pdf")
    assert bad["error"] == "bad_request"
    capped = _ok(
        _call(dirs, "export", dataset_id=ds, sql=sql, format="csv", name="cap", max_rows=1)
    )
    assert capped["truncated"] and capped["rows_returned"] == 1


def test_a_dataset_id_of_another_dir_or_a_bad_id_is_not_found(dirs: Any) -> None:
    assert _call(dirs, "query", dataset_id="ds_" + "0" * 32, sql="SELECT 1")["error"] == "not_found"
    assert _call(dirs, "query", dataset_id="../x", sql="SELECT 1")["error"] == "bad_request"


def test_main_refuses_bad_usage_and_answers_a_missing_request(dirs: Any) -> None:
    assert E.main(["load"], dirs) == 2
    assert E.main(["drop", "r1"], dirs) == 2
    assert E.main(["load", "../r1"], dirs) == 2
    assert E.main(["load", "r1"], dirs) == 0
    answer = json.loads((dirs.run / "answers" / "r1.json").read_text(encoding="utf-8"))
    assert answer == {**answer, "ok": False, "error": "engine_error"}


# --------------------------------------------------------------------------
# Review round 1 (2026-10-06). Each test names its finding.
# --------------------------------------------------------------------------

_FORMULA_SQL = (
    "SELECT 1 AS \"=cmd|' /C calc'!A0\", chr(9) || '=1+1' AS tab, chr(13) || '=2+2' AS cr,"
    " '' AS empty, '-5 apples' AS minus, '@x' AS at, '+1' AS plus, 'plain' AS ok"
)


def test_the_csv_export_guards_every_header_and_every_cell(dirs: Any, invoices: str) -> None:
    """P1-B: a header, an alias and a cell that start with a formula mark get a ' first."""
    _ok(_call(dirs, "export", dataset_id=invoices, sql=_FORMULA_SQL, format="csv", name="g"))
    with open(dirs.outputs / "g.csv", encoding="utf-8", newline="") as fh:
        header, row = list(csv.reader(fh))
    assert header[0] == "'=cmd|' /C calc'!A0" and header[1:] == [
        "tab", "cr", "empty", "minus", "at", "plus", "ok"
    ]  # fmt: skip
    assert row == ["1", "'\t=1+1", "'\r=2+2", "", "'-5 apples", "'@x", "'+1", "plain"]


def test_the_xlsx_export_holds_no_formula_cell(dirs: Any, invoices: str) -> None:
    """P1-B: openpyxl reads each header and each cell back as a value, never a formula."""
    from openpyxl import load_workbook

    head = 'name,"=HYPERLINK(""http://evil.example/"",""x"")",note\nalice,1,=1+1\nbob,2,x\n'
    loaded = _ok(_load(dirs, "h.csv", head))["dataset_id"]
    for ds, sql in ((invoices, _FORMULA_SQL), (loaded, "SELECT * EXCLUDE (_src_row) FROM h")):
        answer = _ok(_call(dirs, "export", dataset_id=ds, sql=sql, format="xlsx", name="f"))
        sheet = load_workbook(dirs.outputs / answer["file"].removeprefix("outputs/"))["export"]
        cells = [c for row in sheet.iter_rows() for c in row]
        assert cells and [c.coordinate for c in cells if c.data_type == "f"] == []
        assert answer["row_ceiling"] == E.XLSX_MAX_ROWS
    assert sheet["B1"].value == '=HYPERLINK("http://evil.example/","x")'


def _thousand(dirs: Any) -> str:
    text = "n,v\n" + "".join(f"{i},{i * 2}\n" for i in range(1000))
    return str(_ok(_load(dirs, "k.csv", text))["dataset_id"])


@pytest.mark.parametrize(
    "sql",
    [
        "SELECT * FROM k",
        "SELECT * FROM k ORDER BY v DESC LIMIT 5",
        "SELECT count(*) FROM k",
        "SELECT max(v) FROM k",
        "SELECT a.n FROM k a JOIN k b ON a.n = b.n",
    ],
)
def test_rows_scanned_counts_each_table_once(dirs: Any, sql: str) -> None:
    """P1-C: a late-materialized scan counts once, and count(*) names its table."""
    answer = _ok(_call(dirs, "query", dataset_id=_thousand(dirs), sql=sql))
    assert answer["rows_scanned"] == {"k": 1000}
    assert [s["table"] for s in answer["source_ranges"]] == ["k"]


def test_preview_and_profile_count_the_whole_table(dirs: Any) -> None:
    ds = _thousand(dirs)
    assert _ok(_call(dirs, "preview", dataset_id=ds, table="k"))["rows_scanned"] == {"k": 1000}
    profile = _ok(_call(dirs, "profile", dataset_id=ds, table="k", column="v"))
    assert profile["rows_scanned"] == {"k": 1000}


def test_a_load_holds_one_row_at_a_time(dirs: Any) -> None:
    """P1-D: the load streams. An 8 MB file costs Python less than half of it.

    The test calls the verb in this process, so ``tracemalloc`` sees it. The
    memory of the image is the Docker half below.
    """
    text = "id,amount,flag,who\n" + "".join(
        f"{i},{i % 997}.5,yes,C{i % 50}\n" for i in range(400_000)
    )
    (dirs.run / "src").mkdir(parents=True)
    (dirs.run / "src" / "big.csv").write_text(text, encoding="utf-8")
    tracemalloc.start()
    try:
        answer = E._load(dirs, {"source": "src/big.csv"}, "rid")
        peak = tracemalloc.get_traced_memory()[1]
    finally:
        tracemalloc.stop()
    assert answer["tables"][0]["rows"] == 400_000
    assert peak < 2**20 * 4, f"{peak / 2**20:.1f} MB for a {len(text) / 2**20:.1f} MB file"


def test_a_load_over_the_row_cap_stops_early(dirs: Any) -> None:
    assert _load(dirs, "r.csv", "a\n" + "1\n" * 50_000, max_rows=10)["error"] == "too_many_rows"


def _serve(dirs: Any, verb: str, **req: Any) -> dict[str, Any]:
    """Run a verb in this process, so that a monkeypatch of the engine applies."""
    return dict(E._serve(dirs, verb, req, uuid.uuid4().hex))


def test_an_export_that_fails_leaves_no_file(dirs: Any, invoices: str, monkeypatch: Any) -> None:
    """P2-F: the time cap, the space cap and an error of another type."""
    slow = "SELECT repeat('x', 1000) AS s FROM range(100000000)"
    answer = _call(dirs, "export", dataset_id=invoices, sql=slow, format="csv", timeout_seconds=1)
    assert answer["error"] == "time" and answer["truncated_reason"] == ["time"]
    monkeypatch.setattr(E.shutil, "disk_usage", lambda _p: type("U", (), {"free": 512})())
    full = _serve(dirs, "export", dataset_id=invoices, sql="SELECT * FROM invoices", format="xlsx")
    assert full["error"] == "too_large" and "MB" in full["message"]
    monkeypatch.undo()

    def boom(*_args: Any) -> Any:
        raise RuntimeError("serialisation failed")

    monkeypatch.setattr(E, "_rows_to", boom)
    broken = _serve(dirs, "export", dataset_id=invoices, sql="SELECT 1", format="xlsx")
    assert broken["error"] == "engine_error" and "serialisation failed" in broken["message"]
    assert not dirs.outputs.exists() or list(dirs.outputs.iterdir()) == []


def _scratch() -> set[str]:
    """The child dirs and the openpyxl spools in the temp dir of this process."""
    root = Path(tempfile.gettempdir())
    return {p.name for pattern in ("data-engine-*", "openpyxl.*") for p in root.glob(pattern)}


def test_a_failed_xlsx_export_leaves_no_spool_and_the_next_verb_works(
    dirs: Any, invoices: str
) -> None:
    """Round 2, P1-1: openpyxl's spool sits in the child's own dir, and it goes.

    The room guard counts each XML escape, so ``&`` costs five bytes.
    """
    before = _scratch()
    amp = "SELECT repeat('&', 2000) AS s FROM range(20000)"
    full = _call(dirs, "export", dataset_id=invoices, sql=amp, format="xlsx", workspace_quota_mb=4)
    assert full["error"] == "too_large", full
    slow = "SELECT range, repeat('y', 50) AS s FROM range(50000000)"
    cut = _call(dirs, "export", dataset_id=invoices, sql=slow, format="xlsx", timeout_seconds=1)
    assert cut["error"] == "time", cut
    assert _scratch() <= before
    assert not dirs.outputs.exists() or list(dirs.outputs.iterdir()) == []
    _ok(_call(dirs, "query", dataset_id=invoices, sql="SELECT 1"))
    assert E._xml_bytes("&<>") == 3 + 4 + 3 + 3 + 48


def test_the_csv_export_stops_at_the_workspace_quota(dirs: Any, invoices: str) -> None:
    """Round 2, P3 (e): ``workspace_quota_mb`` caps a CSV export."""
    big = "SELECT repeat('z', 1000) AS s FROM range(10000)"
    answer = _call(dirs, "export", dataset_id=invoices, sql=big, format="csv", workspace_quota_mb=1)
    assert answer["error"] == "too_large" and "(1 MB)" in answer["message"]
    assert not dirs.outputs.exists() or list(dirs.outputs.iterdir()) == []


def test_an_xlsx_string_is_cut_at_the_excel_cell_limit(dirs: Any, invoices: str) -> None:
    """Round 2, P3 (g)."""
    from openpyxl import load_workbook

    sql = "SELECT repeat('x', 40000) AS s"
    answer = _ok(_call(dirs, "export", dataset_id=invoices, sql=sql, format="xlsx", name="long"))
    assert answer["cells_cut"] == 1 and answer["truncated_reason"] == ["cells"]
    sheet = load_workbook(dirs.outputs / "long.xlsx")["export"]
    assert len(sheet["A2"].value) == E.XLSX_MAX_CELL_CHARS


@pytest.mark.parametrize(
    "text",
    [
        pytest.param('a,b,c\n1,"line1\nline2",10\n2,x,20\n', id="LF"),
        pytest.param('a,b,c\r\n1,"line1\r\nline2",10\r\n2,x,20\r\n', id="CRLF"),
        pytest.param('a,b,c\n1,"line1\rline2",10\n2,x,20\n', id="CR"),
    ],
)
def test_a_quoted_line_break_loads(dirs: Any, text: str) -> None:
    """Round 2, P1-2: no null_padding, so DuckDB reads a quoted line break."""
    ds = _ok(_load(dirs, "q.csv", text))["dataset_id"]
    rows = _ok(_call(dirs, "query", dataset_id=ds, sql="SELECT a, b, c FROM q ORDER BY a"))["rows"]
    assert (
        rows[0][0] == 1 and rows[0][1].replace("\r\n", "\n").replace("\r", "\n") == "line1\nline2"
    )
    assert rows[0][2] == 10 and rows[1] == [2, "x", 20]


def test_short_rows_and_a_wider_row_load(dirs: Any) -> None:
    """Round 2, P1-2: Python pads each row, and a wider row pads the rows above it."""
    ds = _ok(_load(dirs, "s.csv", "a,b,c\n1\n2,y\n3,z,30,extra\n"))["dataset_id"]
    answer = _ok(
        _call(dirs, "query", dataset_id=ds, sql="SELECT * EXCLUDE (_src_row) FROM s ORDER BY a")
    )
    assert [c["name"] for c in answer["columns"]] == ["a", "b", "c", "column_4"]
    assert answer["rows"] == [[1, None, None, None], [2, "y", None, None], [3, "z", 30, "extra"]]


_WIDE_HEADER = ",".join(f"c{i}" for i in range(1000)) + "\n"


@pytest.mark.parametrize(
    "text",
    [
        pytest.param(_WIDE_HEADER + "1\n" * 200_000, id="wide-header"),
        pytest.param("a\n" + "1\n" * 199_999 + ",".join(["9"] * 1000) + "\n", id="late-wide-row"),
    ],
)
def test_a_file_too_wide_for_its_size_is_refused_early(dirs: Any, text: str) -> None:
    """Round 3, P2: the padded stage has a budget of 4 times ``max_file_bytes``.

    With 4 MB here the budget is 16 MB. Padded to 1,000 columns, each file
    would stage about 200 MB.
    """
    started = time.monotonic()
    answer = _load(dirs, "w.csv", text, max_file_bytes=4 * 2**20)
    assert answer["error"] == "bad_file" and "too wide for its size" in answer["message"], answer
    assert time.monotonic() - started < 30
    assert not dirs.data.exists() or list(dirs.data.iterdir()) == []


def test_the_row_cap_names_the_next_step_and_a_huge_row_is_refused(dirs: Any) -> None:
    """Round 2, P3 (b) and (f)."""
    capped = _load(dirs, "r.csv", "a\n" + "1\n" * 50, max_rows=10)
    assert capped["error"] == "too_many_rows" and "Split the file" in capped["message"]
    wide = "a,b,c,d,e,f,g,h,i\n" + ",".join("x" * 120_000 for _ in range(9)) + "\n"
    huge = _load(dirs, "h.csv", wide)
    assert huge["error"] == "bad_file" and "characters" in huge["message"], huge


def test_a_column_of_mixed_date_formats_says_so(dirs: Any) -> None:
    """Round 2, P3 (c)."""
    ds = _ok(_load(dirs, "d.csv", "When\n2026-01-05\n05/01/2026\n2026-02-07\n"))["dataset_id"]
    when = _columns(dirs, ds)["When"]
    assert when["type"] == "text" and when["findings"] == ["mixed_date_formats"]
    assert when["type_counts"] == {"date": 3}


@pytest.mark.parametrize(
    ("fault", "error"),
    [
        pytest.param("exit", "engine_error", id="exit"),
        pytest.param(
            "kill", "memory", id="kill",
            marks=pytest.mark.skipif(os.name != "posix", reason="SIGKILL is POSIX"),
        ),
        pytest.param("sleep", "time", id="sleep"),
    ],
)  # fmt: skip
def test_each_kill_path_of_the_parent_answers(
    dirs: Any, invoices: str, monkeypatch: Any, fault: str, error: str
) -> None:
    """Round 2, P2-5: a child that dies or hangs gets an answer with no rows."""
    monkeypatch.setenv("DATA_ENGINE_FAULT", fault)
    started = time.monotonic()
    answer = _call(dirs, "query", dataset_id=invoices, sql="SELECT 1", timeout_seconds=1)
    assert answer["ok"] is False and answer["error"] == error, answer
    assert not answer.get("rows") and "verb" in answer
    assert time.monotonic() - started < 1 + E.KILL_MARGIN_SECONDS + 10


@pytest.mark.parametrize(
    "sql", ["SELECT [1, 2] AS l", "SELECT uuid() AS u", "SELECT '\\x00'::BLOB AS b"]
)
def test_an_xlsx_export_of_an_odd_type_writes_its_text(dirs: Any, invoices: str, sql: str) -> None:
    """P2-F: a list, a UUID and a BLOB become text, and the file opens."""
    from openpyxl import load_workbook

    answer = _ok(_call(dirs, "export", dataset_id=invoices, sql=sql, format="xlsx", name="odd"))
    sheet = load_workbook(dirs.outputs / "odd.xlsx")["export"]
    assert isinstance(sheet["A2"].value, str)
    assert answer["rows_returned"] == 1
    assert [p.name for p in dirs.outputs.iterdir()] == ["odd.xlsx"]


@pytest.mark.parametrize("pad", ["é", "€", "😀"])
@pytest.mark.parametrize(
    "stmt", ["SHOW SELECT 1", "DESCRIBE SELECT 1", "SUMMARIZE SELECT 1", "PRAGMA version"]
)
def test_a_multibyte_comment_does_not_hide_the_first_word(pad: str, stmt: str) -> None:
    """P3 (a): the tokenizer gives byte offsets."""
    for k in range(1, 21):
        with pytest.raises(E.Refused):
            E._check_sql("/*" + pad * k + "*/" + stmt)
    assert E._check_sql("/*" + pad * 7 + "*/ SELECT 1") == "/*" + pad * 7 + "*/ SELECT 1"


def test_a_long_run_of_spaces_loads_quickly(dirs: Any) -> None:
    """P3 (b): the number patterns are RE2, so a long value costs linear time."""
    started = time.monotonic()
    answer = _ok(_load(dirs, "w.csv", "a,b\n1" + " " * 20_000 + "x,1\n2,2\n"))
    assert time.monotonic() - started < 10
    assert _columns(dirs, answer["dataset_id"])["a"]["type"] == "text"


def test_the_load_reads_its_clock(dirs: Any) -> None:
    """P3 (b): the record loop stops at the time cap."""
    path = dirs.run / "t.csv"
    path.parent.mkdir(parents=True)
    path.write_text("a\n" * 5000, encoding="utf-8")
    with pytest.raises(E.Refused, match="time cap"):
        list(E._records(path, ",", deadline=0))


def test_a_long_header_keeps_each_answer_under_1_mb(dirs: Any) -> None:
    """P3 (c): an answer cuts each name, and the manifest keeps it whole."""
    head = ",".join(f"h{i}_" + "x" * 5000 for i in range(400))
    answer = _ok(_load(dirs, "l.csv", head + "\n" + ",".join("1" for _ in range(400)) + "\n"))
    ds = answer["dataset_id"]
    query = _ok(_call(dirs, "query", dataset_id=ds, sql="SELECT * FROM l"))
    sizes = [p.stat().st_size for p in (dirs.run / "answers").glob("*.json")]
    assert max(sizes) < E.MAX_ANSWER_BYTES
    assert max(len(c["name"]) for c in answer["tables"][0]["columns"]) == E.MAX_NAME_CHARS + 1
    assert max(len(c["name"]) for c in query["columns"]) == E.MAX_NAME_CHARS + 1
    assert len(_columns(dirs, ds)["h0_" + "x" * 5000]["name"]) == 5003


@pytest.mark.parametrize("kind", ["csv", "xlsx"])
def test_each_statement_of_a_load_runs_on_its_own_database(
    dirs: Any, monkeypatch: Any, kind: str
) -> None:
    """H-290: no DuckDB database runs two statements of a load.

    One database keeps memory from each wide statement, outside its
    ``memory_limit``. Over the 41 statements of a load of 400 columns it grew
    to 500 MB resident in the image, and to 771 MB on a CI runner, past the
    watch at 768 MB. So a wide load failed at random with ``error: memory``.
    The Docker half below measures the peak. This test holds the cause on
    every host, Windows too.
    """
    real_connect, real_timed = E._load_connection, E._timed
    opened: list[Any] = []
    ran: list[Any] = []

    def connect(stage: Path) -> Any:
        opened.append(real_connect(stage))
        return opened[-1]

    def timed(con: Any, timeout: int, work: Any, what: str = "query") -> Any:
        ran.append(con)
        return real_timed(con, timeout, work, what)

    monkeypatch.setattr(E, "_load_connection", connect)
    monkeypatch.setattr(E, "_timed", timed)
    width = 70  # Three batches of columns, so one shared database would run many statements.
    cells = [
        [str(r * j % 97) + (".5" if j % 3 == 0 else "") for j in range(width)] for r in range(5)
    ]
    (dirs.run / "src").mkdir(parents=True)
    if kind == "csv":
        lines = [",".join(f"h{j}" for j in range(width))] + [",".join(r) for r in cells]
        (dirs.run / "src" / "w.csv").write_text("\n".join(lines) + "\n", encoding="utf-8")
    else:
        head = "".join(inline(f"{E._letter(j)}1", f"h{j}") for j in range(width))
        rows = xrow(1, head) + "".join(
            xrow(i, "".join(num(f"{E._letter(j)}{i}", v) for j, v in enumerate(r)))
            for i, r in enumerate(cells, 2)
        )
        (dirs.run / "src" / "w.xlsx").write_bytes(_book([("Data", xb.sheet(rows))]))
    answer = E._serve(dirs, "load", {"source": f"src/w.{kind}"}, "rid")
    assert answer["ok"] and len(answer["tables"][0]["columns"]) == width, answer
    assert len(ran) > 10, len(ran)
    assert len({id(con) for con in ran}) == len(ran), "a database ran two statements of a load"
    for con in opened:
        with pytest.raises(duckdb.ConnectionException):
            con.execute("SELECT 1")


def test_text_dates_and_the_thousands_finding_and_a_huge_power(dirs: Any) -> None:
    """P3 (d), (e) and (f)."""
    mixed = _ok(_load(dirs, "m.csv", 'Code\n1\nx\n2026-01-05\nyes\n"Jan 5, 2026"\n'))
    assert _columns(dirs, mixed["dataset_id"])["Code"]["type_counts"] == {
        "number": 1, "boolean": 1, "date": 2, "text": 1
    }  # fmt: skip
    dotted = _ok(_load(dirs, "d.csv", "Qty;Name\n1.000;a\n2.500;b\n"))
    qty = _columns(dirs, dotted["dataset_id"])["Qty"]
    assert qty["type"] == "decimal" and qty["findings"] == ["ambiguous_thousands"]
    huge = _ok(_load(dirs, "e.csv", "Big,Small\n1e400,1e5\n2,3\n"))
    cols = _columns(dirs, huge["dataset_id"])
    assert cols["Big"]["type"] == "text" and cols["Small"]["type"] == "decimal"


def test_a_nul_byte_reads_as_a_replacement_mark(dirs: Any) -> None:
    answer = _ok(_load(dirs, "n.csv", b"a,b\nx\x00y,1\nz,2\n"))
    assert _columns(dirs, answer["dataset_id"])["a"]["samples"] == ["x�y", "z"]


# --------------------------------------------------------------------------
# WS-43y1b: the .xlsx reader, the nine layout rules of §7.10, Excel dates, and
# the zip and XML checks. Each workbook is built in code with zipfile, by the
# builder of EM-T11b (tests/unit/_xlsx_build.py), so a test can make any part
# hostile. One test reads a workbook that openpyxl wrote.
# --------------------------------------------------------------------------

#: Cell styles: 0 General, 1 a date (14), 2 a date and time (164), 3 a time
#: (20) and 4 a duration (46).
_STYLES = (
    f'<?xml version="1.0" encoding="UTF-8"?><styleSheet xmlns="{xb.MAIN}">'
    '<numFmts count="1"><numFmt numFmtId="164" formatCode="dd/mm/yyyy hh:mm"/></numFmts>'
    '<cellXfs count="5"><xf numFmtId="0"/><xf numFmtId="14"/><xf numFmtId="164"/>'
    '<xf numFmtId="20"/><xf numFmtId="46"/></cellXfs></styleSheet>'
).encode()


def _f(ref: str, formula: str, cached: object | None = None, kind: str = "") -> str:
    """A formula cell, with its cached value when one is given."""
    t = f' t="{kind}"' if kind else ""
    v = f"<v>{cached}</v>" if cached is not None else ""
    return f'<c r="{ref}"{t}><f>{formula}</f>{v}</c>'


def _date(ref: str, serial: object, style: int = 1) -> str:
    return f'<c r="{ref}" s="{style}"><v>{serial}</v></c>'


def _book(
    sheets: list[tuple[str, bytes]],
    *,
    hidden: frozenset[str] = frozenset(),
    date1904: bool = False,
    extra: dict[str, bytes] | None = None,
) -> bytes:
    """A workbook of *sheets*, with ``_STYLES`` at the usual name."""
    files = xb.parts(sheets, hidden=hidden)
    if date1904:
        files["xl/workbook.xml"] = files["xl/workbook.xml"].replace(
            b"<sheets>", b'<workbookPr date1904="1"/><sheets>'
        )
    files["xl/styles.xml"] = _STYLES
    files.update(extra or {})
    return xb.package(files)


def _rows(*rows: tuple[int, str]) -> str:
    return "".join(xrow(n, cells) for n, cells in rows)


def _tables(dirs: Any, ds_id: str) -> dict[str, dict[str, Any]]:
    manifest = json.loads((dirs.data / ds_id / "manifest.json").read_text(encoding="utf-8"))
    return {t["name"]: t for t in manifest["tables"]}


def _sql(dirs: Any, ds_id: str, sql: str) -> list[list[Any]]:
    return list(_ok(_call(dirs, "query", dataset_id=ds_id, sql=sql))["rows"])


#: Rules 1 to 6: a title row, a header of two rows under two merged cells, a
#: total row named "Total", and a note under the table after a gap.
_REPORT = _rows(
    (1, inline("A1", "Sales report 2026")),
    (3, inline("A3", "Region") + inline("B3", "Sales") + inline("D3", "Cost")),
    (4, inline("B4", "Q1") + inline("C4", "Q2") + inline("D4", "Q1")),
    (5, inline("A5", "North") + num("B5", 10) + num("C5", 20) + num("D5", 5)),
    (6, inline("A6", "South") + num("B6", 30) + num("C6", 40) + num("D6", 6)),
    (7, inline("A7", "Total") + _f("B7", "SUM(B5:B6)", 40) + num("C7", 60) + num("D7", 11)),
    (9, inline("A9", "Source: finance")),
)
_REPORT_MERGES = (
    '<mergeCells count="3"><mergeCell ref="A1:D1"/><mergeCell ref="A3:A4"/>'
    '<mergeCell ref="B3:C3"/></mergeCells>'
)


def test_a_workbook_follows_the_layout_rules(dirs: Any) -> None:
    """Rules 1 to 6 of §7.10, and a cell reference of a sheet name with a space."""
    data = _book([("Q1 Sales", xb.sheet(_REPORT, after=_REPORT_MERGES))])
    answer = _ok(_load(dirs, "r.xlsx", data))
    ds = answer["dataset_id"]
    table = _tables(dirs, ds)["q1_sales"]
    assert [c["name"] for c in table["columns"]] == [
        "Region",
        "Sales / Q1",
        "Sales / Q2",
        "Cost / Q1",
    ]
    assert table["rows"] == 2 and table["range"] == "'Q1 Sales'!A5:D6"
    assert table["left_out"] == [
        {"what": "title rows", "range": "'Q1 Sales'!A1:A1"},
        {"what": "header rows", "range": "'Q1 Sales'!A3:D4"},
        {"what": "notes", "range": "'Q1 Sales'!A9:A9"},
        {"what": "total rows", "table": "q1_sales__totals", "count": 1, "rows": [7]},
    ]
    # Rule 5: the sum never counts the total row, and the side table keeps it.
    assert _sql(dirs, ds, 'SELECT sum("Sales / Q1"), sum("Cost / Q1") FROM q1_sales') == [[40, 11]]
    assert _sql(dirs, ds, 'SELECT "Sales / Q1" FROM q1_sales__totals') == [[40]]
    preview = _ok(_call(dirs, "preview", dataset_id=ds, table="q1_sales"))
    assert preview["cells"] == ["'Q1 Sales'!A5:D5", "'Q1 Sales'!A6:D6"]
    assert answer["sheets"] == [
        {"name": "Q1 Sales", "hidden": False, "tables": ["q1_sales"], "left_out": []}
    ]


def test_a_row_of_column_sums_is_a_total_row(dirs: Any) -> None:
    """Rule 4 with no word: a last row of sums with no text is a total row. A
    row that equals a sum and has a data row after it is data, and so is a
    row of sums with text in it."""
    head = inline("A1", "Item") + inline("B1", "Qty") + inline("C1", "Amount")

    def item(n: int, name: str | None, qty: int, amount: int) -> tuple[int, str]:
        label = inline(f"A{n}", name) if name else ""
        return n, label + num(f"B{n}", qty) + num(f"C{n}", amount)

    sums = _rows((1, head), item(2, "a", 1, 10), item(3, "b", 1, 20), item(4, "c", 2, 30),
                 (5, _f("B5", "SUM(B2:B4)", 4) + _f("C5", "SUM(C2:C4)", 60)))  # fmt: skip
    mid = _rows((1, head), item(2, "x", 1, 1), item(3, "y", 1, 1), item(4, None, 2, 2),
                item(5, "z", 5, 5))  # fmt: skip
    ds = _ok(_load(dirs, "s.xlsx", _book([("Sums", xb.sheet(sums)), ("Mid", xb.sheet(mid))])))[
        "dataset_id"
    ]
    tables = _tables(dirs, ds)
    assert tables["sums"]["rows"] == 3 and tables["sums__totals"]["rows"] == 1
    assert _sql(dirs, ds, "SELECT sum(Amount), sum(Qty) FROM sums") == [[60, 4]]
    assert _sql(dirs, ds, "SELECT Amount FROM sums__totals") == [[60]]
    assert tables["mid"]["rows"] == 4 and "mid__totals" not in tables
    assert _sql(dirs, ds, "SELECT sum(Amount) FROM mid") == [[9]]


def test_a_formula_gives_its_cached_value_and_never_its_text(dirs: Any) -> None:
    """Rule 8. A formula with no cached value, or an error value, reads as
    empty, and its column counts it."""
    sheet = _rows(
        (1, inline("A1", "Qty") + inline("B1", "Double") + inline("C1", "Tag") + inline("D1", "Now")),
        (2, num("A2", 3) + _f("B2", "A2*2", 6) + _f("C2", '"x"&amp;"y"', "xy", "str")
         + _f("D2", "1/0", "#DIV/0!", "e")),
        (3, num("A3", 4) + _f("B3", "A3*2", 8) + _f("C3", '"z"', "z", "str") + _f("D3", "NOW()")),
    )  # fmt: skip
    ds = _ok(_load(dirs, "f.xlsx", _book([("F", xb.sheet(sheet))])))["dataset_id"]
    rows = _sql(dirs, ds, "SELECT Qty, Double, Tag, Now FROM f ORDER BY Qty")
    assert rows == [[3, 6, "xy", None], [4, 8, "z", None]]
    cols = {c["name"]: c for c in _tables(dirs, ds)["f"]["columns"]}
    assert cols["Double"]["type"] == "integer"
    assert cols["Now"]["uncached_formulas"] == 1 and cols["Now"]["error_values"] == 1
    assert "A2*2" not in (dirs.data / ds / "manifest.json").read_text(encoding="utf-8")


def test_several_tables_on_one_sheet_and_a_blank_separator_row(dirs: Any) -> None:
    """Rule 7: an empty row before a new header, or an empty column, splits a
    sheet. An empty row before a data row splits no table."""
    two = _rows(
        (1, inline("A1", "Name") + inline("B1", "Qty") + inline("E1", "Code") + inline("F1", "Value")),
        (2, inline("A2", "a") + num("B2", 1) + inline("E2", "k1") + num("F2", 100)),
        (3, inline("A3", "b") + num("B3", 2) + inline("E3", "k2") + num("F3", 200)),
        (5, inline("A5", "City") + inline("B5", "Pop")),
        (6, inline("A6", "Pune") + num("B6", 7)),
        (7, inline("A7", "Goa") + num("B7", 2)),
    )  # fmt: skip
    gap = _rows(
        (1, inline("A1", "Name") + inline("B1", "Qty")),
        (2, inline("A2", "a") + num("B2", 1)),
        (3, inline("A3", "b") + num("B3", 2)),
        (5, inline("A5", "c") + num("B5", 3)),
    )
    ds = _ok(_load(dirs, "t.xlsx", _book([("Two", xb.sheet(two)), ("Gap", xb.sheet(gap))])))[
        "dataset_id"
    ]
    tables = _tables(dirs, ds)
    ranges = {t["range"]: [c["name"] for c in t["columns"]] for t in tables.values()}
    assert ranges == {
        "Two!A2:B3": ["Name", "Qty"], "Two!A6:B7": ["City", "Pop"],
        "Two!E2:F3": ["Code", "Value"], "Gap!A2:B5": ["Name", "Qty"],
    }  # fmt: skip
    assert tables["gap"]["rows"] == 3
    assert {"what": "empty rows", "count": 1, "rows": [4]} in tables["gap"]["left_out"]
    side = next(t for t in tables.values() if t["range"] == "Two!E2:F3")
    rows = _ok(_call(dirs, "preview", dataset_id=ds, table=side["name"]))["cells"]
    assert rows == ["Two!E2:F2", "Two!E3:F3"]


def test_hidden_sheets_rows_and_columns_are_read_and_marked(dirs: Any) -> None:
    """Rule 9."""
    cols = '<cols><col min="2" max="2" width="0" hidden="1"/></cols>'
    open_sheet = (
        '<row r="1"><c r="A1" t="inlineStr"><is><t>Item</t></is></c>'
        '<c r="B1" t="inlineStr"><is><t>Cost</t></is></c></row>'
        '<row r="2"><c r="A2" t="inlineStr"><is><t>a</t></is></c><c r="B2"><v>5</v></c></row>'
        '<row r="3" hidden="1"><c r="A3" t="inlineStr"><is><t>b</t></is></c><c r="B3"><v>7</v></c></row>'
    )
    rates = _rows(
        (1, inline("A1", "Rate") + inline("B1", "Code")), (2, num("A2", 9) + inline("B2", "x"))
    )
    data = _book([("Open", xb.sheet(open_sheet, before=cols)), ("Rates", xb.sheet(rates))],
                 hidden=frozenset({"Rates"}))  # fmt: skip
    ds = _ok(_load(dirs, "h.xlsx", data))["dataset_id"]
    tables = _tables(dirs, ds)
    assert tables["rates"]["hidden"] is True and tables["open"]["hidden"] is False
    assert tables["open"]["hidden_rows"] == {"count": 1, "rows": [3]}
    assert [c.get("hidden", False) for c in tables["open"]["columns"]] == [False, True]
    assert _sql(dirs, ds, "SELECT sum(Cost) FROM open") == [[12]]


@pytest.mark.parametrize(
    ("serial", "kind", "date1904", "text"),
    [
        (1, "date", False, "1900-01-01"),
        (59, "date", False, "1900-02-28"),
        (60, "date", False, "1900-02-29"),  # Excel's leap-year bug: a day that did not exist
        (61, "date", False, "1900-03-01"),
        (45658, "date", False, "2025-01-01"),
        (45658.5, "date", False, "2025-01-01 12:00:00"),  # a date format keeps its time
        (45658.25, "datetime", False, "2025-01-01 06:00:00"),
        (45658, "datetime", False, "2025-01-01 00:00:00"),
        (0.4375, "time", False, "10:30:00"),
        (0, "date", True, "1904-01-01"),
        (44196, "date", True, "2025-01-01"),
        (-1, "date", False, None),
        (2_958_466, "date", False, None),
    ],
)  # fmt: skip
def test_an_excel_serial_date(serial: float, kind: str, date1904: bool, text: str | None) -> None:
    assert E._excel_date(serial, kind, date1904) == text


def test_excel_dates_load_with_their_types(dirs: Any) -> None:
    """A cell with a date format is a date. A number with no date format stays
    a number, and a date header gives it the finding of a serial date."""
    head = "".join(inline(f"{c}1", n) for c, n in zip("ABCDEF", ["When", "At", "Clock", "Spent",
                                                                  "Order date", "Day"], strict=True))  # fmt: skip
    rows = _rows(
        (1, head),
        (2, _date("A2", 45658) + _date("B2", 45658.25, 2) + _date("C2", 0.5, 3)
         + _date("D2", 1.5, 4) + num("E2", 45700) + '<c r="F2" t="d"><v>2026-01-05</v></c>'),
        (3, _date("A3", 45700) + _date("B3", 45700, 2) + _date("C3", 0.25, 3)
         + _date("D3", 2, 4) + num("E3", 45701) + '<c r="F3" t="d"><v>2026-01-06</v></c>'),
    )  # fmt: skip
    ds = _ok(_load(dirs, "d.xlsx", _book([("D", xb.sheet(rows))])))["dataset_id"]
    cols = {c["name"]: c for c in _tables(dirs, ds)["d"]["columns"]}
    types = {name: c["type"] for name, c in cols.items()}
    assert types == {"When": "date", "At": "datetime", "Clock": "text", "Spent": "decimal",
                     "Order date": "integer", "Day": "date"}  # fmt: skip
    assert cols["When"]["min"] == "2025-01-01" and cols["At"]["min"] == "2025-01-01T06:00:00"
    assert cols["Clock"]["samples"] == ["12:00:00", "06:00:00"]
    assert cols["Order date"]["findings"] == ["possible_excel_serial_date"]
    old = _book([("D", xb.sheet(_rows((1, inline("A1", "When")), (2, _date("A2", 44196)))))],
                date1904=True)  # fmt: skip
    cols = _columns(dirs, _ok(_load(dirs, "o.xlsx", old))["dataset_id"])
    assert cols["When"]["type"] == "date" and cols["When"]["min"] == "2025-01-01"


def test_the_request_names_the_header_rows_and_the_sheet(dirs: Any) -> None:
    two = _rows(
        (1, inline("A1", "Sales") + inline("B1", "Sales")),
        (2, inline("A2", "Q1") + inline("B2", "Q2")),
        (3, num("A3", 1) + num("B3", 2)),
    )
    other = _rows((1, inline("A1", "x") + inline("B1", "y")), (2, num("A2", 1) + num("B2", 2)))
    data = _book([("First", xb.sheet(two)), ("Second", xb.sheet(other))])
    auto = _ok(_load(dirs, "a.xlsx", data))
    assert [c["name"] for c in auto["tables"][0]["columns"]] == ["Sales", "Sales_2"]
    named = _ok(_load(dirs, "a.xlsx", data, header_rows=2, sheet="first"))
    assert [t["name"] for t in named["tables"]] == ["first"]
    assert [c["name"] for c in named["tables"][0]["columns"]] == ["Sales / Q1", "Sales / Q2"]
    assert named["dataset_id"] != auto["dataset_id"]
    missing = _load(dirs, "a.xlsx", data, sheet="Third")
    assert missing["error"] == "not_found" and "First, Second" in missing["message"]
    assert _load(dirs, "c.csv", "a\n1\n", sheet="First")["error"] == "bad_request"
    many = _book([(f"S{i}", xb.sheet(xrow(1, inline("A1", "x")))) for i in range(E.MAX_SHEETS + 1)])
    assert _load(dirs, "m.xlsx", many)["error"] == "too_many_sheets"


def test_a_column_of_text_rows_and_a_sheet_of_titles(dirs: Any) -> None:
    """With no wider row, the rows of one cell each are a table of one column.
    A workbook with no data row under a header is refused."""
    names = _rows((1, inline("A1", "Name")), (2, inline("A2", "Alice")), (3, inline("A3", "Bob")))
    ds = _ok(_load(dirs, "n.xlsx", _book([("N", xb.sheet(names))])))["dataset_id"]
    assert _sql(dirs, ds, "SELECT Name FROM n ORDER BY Name") == [["Alice"], ["Bob"]]
    title = _book([("T", xb.sheet(_rows((1, inline("A1", "Only a title")))))])
    assert _load(dirs, "t.xlsx", title)["error"] == "empty"


def test_an_xlsm_file_loads_its_values_and_runs_no_macro(dirs: Any) -> None:
    data = _book([("Q1 Sales", xb.sheet(_REPORT, after=_REPORT_MERGES))],
                 extra={"xl/vbaProject.bin": b"\xd0\xcf\x11\xe0 macro bytes"})  # fmt: skip
    ds = _ok(_load(dirs, "r.xlsm", data))["dataset_id"]
    assert _sql(dirs, ds, 'SELECT sum("Sales / Q1") FROM q1_sales') == [[40]]


def test_a_workbook_that_openpyxl_wrote_loads(dirs: Any, tmp_path: Path) -> None:
    """A real writer: shared strings, styles, a merged header and a formula
    with no cached value, which openpyxl never writes."""
    from datetime import date as day

    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Orders"
    sheet.append(["Orders of March"])
    sheet.append([])
    sheet.append(["Order", "Placed", "Amount", "Tax"])
    sheet.append(["O-1", day(2026, 3, 1), 120.5, "=C4*0.18"])
    sheet.append(["O-2", day(2026, 3, 9), 80, "=C5*0.18"])
    sheet.merge_cells("A1:D1")
    book.create_sheet("Notes").sheet_state = "hidden"
    path = tmp_path / "o.xlsx"
    book.save(path)
    ds = _ok(_load(dirs, "o.xlsx", path.read_bytes()))["dataset_id"]
    cols = _columns(dirs, ds)
    assert cols["Placed"]["type"] == "date" and cols["Placed"]["max"] == "2026-03-09"
    assert cols["Amount"]["type"] == "decimal" and cols["Tax"]["uncached_formulas"] == 2
    assert _sql(dirs, ds, "SELECT sum(Amount) FROM orders") == [["200.5"]]


def _python_peak(dirs: Any, rows: int) -> int:
    """The Python peak (``tracemalloc``) of a load of *rows* rows, each with its own shared string."""
    parts = xb.parts([("Big", b"")], shared=b"")
    strings = "".join(f"<si><t>Customer {i:06d}</t></si>" for i in range(rows))
    parts["xl/sharedStrings.xml"] = f'<sst xmlns="{xb.MAIN}">{strings}</sst>'.encode()
    body = "".join(
        xrow(i, num(f"A{i}", i) + f'<c r="B{i}" t="s"><v>{i - 2}</v></c>' + num(f"C{i}", f"{i}.5"))
        for i in range(2, rows + 2)
    )
    head = xrow(1, inline("A1", "id") + inline("B1", "who") + inline("C1", "amount"))
    parts["xl/worksheets/sheet1.xml"] = xb.sheet(head + body)
    (dirs.run / "src").mkdir(parents=True, exist_ok=True)
    (dirs.run / "src" / f"big{rows}.xlsx").write_bytes(xb.package(parts))
    del parts, body, strings
    tracemalloc.start()
    try:
        answer = E._load(dirs, {"source": f"src/big{rows}.xlsx"}, f"rid{rows}")
        peak = tracemalloc.get_traced_memory()[1]
    finally:
        tracemalloc.stop()
    assert answer["tables"][0]["rows"] == rows
    return int(peak)


def test_a_workbook_load_streams(dirs: Any, monkeypatch: Any) -> None:
    """The load holds one row at a time, and the shared strings are packed.

    A chunk of 64 KB keeps the buffers of the parse small, so the peak shows
    what grows with the file. 60,000 more rows, each with its own shared
    string, cost Python less than 2 MB more: about 17 bytes a string. A list
    of the strings would cost about 4 MB, and a list of the rows far more.
    The memory of the image is the Docker half below.
    """
    monkeypatch.setattr(E, "_XML_CHUNK", 2**16)
    small, large = _python_peak(dirs, 20_000), _python_peak(dirs, 80_000)
    assert small < 4 * 2**20, f"{small / 2**20:.1f} MB"
    assert large - small < 2 * 2**20, f"{small / 2**20:.1f} MB, then {large / 2**20:.1f} MB"


# The zip and XML checks (§7.10 "Isolation"). Each refusal comes before any
# part is parsed, or at the first event of the part.


def _stream_part(files: dict[str, bytes], name: str, size: int, fill: bytes = b"A") -> bytes:
    """A package whose part *name* holds *size* bytes of *fill* in one inline string."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for part, data in files.items():
            if part != name:
                zf.writestr(part, data)
        with zf.open(name, "w", force_zip64=True) as fh:
            fh.write(
                f'<worksheet xmlns="{xb.MAIN}"><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>'.encode()
            )
            chunk = fill * 2**20
            for _ in range(size // 2**20):
                fh.write(chunk)
            fh.write(b"</t></is></c></row></sheetData></worksheet>")
    return buf.getvalue()


def _spy_open(monkeypatch: Any) -> list[str]:
    """Record each part that the engine unpacks."""
    opened: list[str] = []
    real = zipfile.ZipFile.open

    def spy(self: Any, name: Any, *args: Any, **kwargs: Any) -> Any:
        opened.append(getattr(name, "filename", name))
        return real(self, name, *args, **kwargs)

    monkeypatch.setattr(zipfile.ZipFile, "open", spy)
    return opened


def _serve_book(dirs: Any, data: bytes, name: str = "b.xlsx") -> dict[str, Any]:
    (dirs.run / "src").mkdir(parents=True, exist_ok=True)
    (dirs.run / "src" / name).write_bytes(data)
    return _serve(dirs, "load", source=f"src/{name}", max_file_bytes=50 * 2**20)


def test_a_zip_bomb_is_refused_before_any_part_is_unpacked(dirs: Any, monkeypatch: Any) -> None:
    """A part over 200 MB unpacked, and a part over 100 times its packed size.

    The sizes are numbers, and not the engine's caps, so a changed cap fails
    here and builds no larger bomb.
    """
    files = xb.parts([("S", b"")])
    big = _stream_part(files, "xl/worksheets/sheet1.xml", 201 * 2**20)
    dense = _stream_part(files, "xl/worksheets/sheet1.xml", 8 * 2**20)
    assert len(big) < 2**20 and len(dense) < 2**20
    opened = _spy_open(monkeypatch)
    answer = _serve_book(dirs, big)
    assert answer["error"] == "too_large" and "200 MB" in answer["message"], answer
    ratio = _serve_book(dirs, dense, "d.xlsx")
    assert ratio["error"] == "bad_file" and "100 times" in ratio["message"], ratio
    assert opened == []


def test_a_part_that_hides_its_size_is_refused(dirs: Any) -> None:
    """zipfile stops at the declared size and checks the CRC there."""
    from tests.unit.test_attachment_xlsx import _lie_about_size

    data = _book([("Q1 Sales", xb.sheet(_REPORT, after=_REPORT_MERGES))])
    answer = _serve_book(dirs, _lie_about_size(data, b"xl/worksheets/sheet1.xml", 60))
    assert answer["error"] == "bad_file", answer


def test_too_many_zip_entries_are_refused_before_the_directory_is_read(
    dirs: Any, monkeypatch: Any
) -> None:
    files = xb.parts([("S", xb.sheet(_REPORT))])
    files |= {f"docProps/x{i}.xml": b"" for i in range(E.ZIP_MAX_ENTRIES)}
    data = xb.package(files)
    made: list[Any] = []
    monkeypatch.setattr(E.zipfile, "ZipFile", lambda *a, **k: made.append(a))
    answer = _serve_book(dirs, data)
    assert answer["error"] == "bad_file" and "zip entries" in answer["message"], answer
    assert made == []


def test_the_entry_count_reads_a_zip64_end_record(tmp_path: Path) -> None:
    """Past 65,535 entries, the end record says 0xFFFF, and the zip64 end
    record holds the count. A file with no end record is no workbook."""
    path = tmp_path / "many.zip"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_STORED) as zf:
        for i in range(65_536):
            zf.writestr(f"p{i}", b"")
    assert E._zip_entries(path) == 65_536
    plain = tmp_path / "plain.zip"
    plain.write_bytes(xb.package({"a": b"1", "b": b"2"}))
    assert E._zip_entries(plain) == 2
    (tmp_path / "x.xlsx").write_bytes(b"not a zip at all")
    with pytest.raises(E.Refused):
        E._zip_entries(tmp_path / "x.xlsx")


_LAUGHS = (
    '<?xml version="1.0"?><!DOCTYPE lolz [<!ENTITY lol "lol">'
    '<!ENTITY lol2 "&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;">]>'
)


@pytest.mark.parametrize(
    ("part", "body", "says"),
    [
        pytest.param("xl/worksheets/sheet1.xml",
                     _LAUGHS + f'<worksheet xmlns="{xb.MAIN}"><sheetData/></worksheet>', "DTD",
                     id="sheet-billion-laughs"),
        pytest.param("xl/sharedStrings.xml",
                     _LAUGHS + f'<sst xmlns="{xb.MAIN}"><si><t>&lol2;</t></si></sst>', "DTD",
                     id="strings-entity"),
        pytest.param("xl/workbook.xml",
                     '<?xml version="1.0"?><!DOCTYPE w SYSTEM "http://x/w.dtd"><workbook/>', "DTD",
                     id="workbook-external-dtd"),
        pytest.param("xl/styles.xml",
                     '<!DOCTYPE s [<!ENTITY % p SYSTEM "file:///etc/passwd"> %p;]><styleSheet/>',
                     "DTD", id="styles-parameter-entity"),
        pytest.param("xl/worksheets/sheet1.xml",
                     "<a>" * 100 + "</a>" * 100, "too deeply", id="deep-nesting"),
        pytest.param("xl/worksheets/sheet1.xml",
                     '<?xml version="1.0" encoding="UTF-16"?><worksheet/>'.encode("utf-16"),
                     "readable", id="utf-16"),
    ],
)  # fmt: skip
def test_an_xml_bomb_is_refused_at_its_first_event(
    dirs: Any, part: str, body: str | bytes, says: str
) -> None:
    files = xb.parts([("S", xb.sheet(_REPORT))], shared=xb.strings(["x"]))
    files["xl/styles.xml"] = _STYLES
    files[part] = body.encode() if isinstance(body, str) else body
    answer = _serve_book(dirs, xb.package(files))
    assert answer["error"] == "bad_file" and says in answer["message"], answer


@pytest.mark.parametrize(
    ("rels_part", "target"),
    [("xl/_rels/workbook.xml.rels", "../../evil.xml"), ("_rels/.rels", "docs/workbook.xml")],
)
def test_a_part_outside_xl_is_refused(dirs: Any, rels_part: str, target: str) -> None:
    files = xb.parts([("S", xb.sheet(_REPORT))])
    kind = xb.WORKSHEET if rels_part.startswith("xl/") else xb.OFFICE_DOCUMENT
    files[rels_part] = xb.rels([("rId1", kind, target)])
    files[target.lstrip("./")] = xb.sheet(_REPORT)
    answer = _serve_book(dirs, xb.package(files))
    assert answer["error"] == "bad_file" and "outside xl/" in answer["message"], answer


def test_a_password_a_broken_zip_and_rows_out_of_order_are_refused(dirs: Any) -> None:
    good = _book([("S", xb.sheet(_REPORT))])
    locked = bytearray(good)
    at = locked.find(b"PK\x01\x02")
    while at != -1:  # Set the "encrypted" bit of each entry in the directory.
        locked[at + 8] |= 1
        at = locked.find(b"PK\x01\x02", at + 4)
    assert _serve_book(dirs, bytes(locked))["error"] == "bad_file"
    assert (
        _serve_book(dirs, b"\xd0\xcf\x11\xe0 an OLE file, as Excel encrypts")["error"] == "bad_file"
    )
    backwards = _rows((5, inline("A5", "a") + inline("B5", "b")), (3, num("A3", 1) + num("B3", 2)))
    answer = _serve_book(dirs, _book([("S", xb.sheet(backwards))]))
    assert answer["error"] == "bad_file" and "out of order" in answer["message"]


def test_the_caps_of_a_workbook_hold(dirs: Any) -> None:
    """The rows of a dataset, the width of a table and the padded stage (round 3)."""
    rows = [(1, inline("A1", "a") + inline("B1", "b"))]
    rows += [(i, num(f"A{i}", i) + num(f"B{i}", i)) for i in range(2, 30)]
    (dirs.run / "src").mkdir(parents=True)
    (dirs.run / "src" / "r.xlsx").write_bytes(_book([("R", xb.sheet(_rows(*rows)))]))
    capped = _serve(dirs, "load", source="src/r.xlsx", max_rows=10)
    assert capped["error"] == "too_many_rows"
    wide_head = "".join(inline(f"{E._letter(j)}1", f"h{j}") for j in range(E.MAX_COLUMNS + 1))
    wide = _serve_book(dirs, _book([("W", xb.sheet(xrow(1, wide_head) + xrow(2, num("A2", 1))))]))
    assert wide["error"] == "too_many_columns"
    head = "".join(inline(f"{E._letter(j)}1", f"h{j}") for j in range(900))
    tall = xrow(1, head) + "".join(xrow(i, num(f"A{i}", 1)) for i in range(2, 6000))
    (dirs.run / "src" / "p.xlsx").write_bytes(_book([("P", xb.sheet(tall))]))
    padded = _serve(dirs, "load", source="src/p.xlsx", max_file_bytes=2**20)
    assert padded["error"] == "bad_file" and "too wide for its size" in padded["message"]


def test_a_long_sheet_name_is_cut_to_31_characters_and_stays_unique(dirs: Any) -> None:
    """Security review P3-1. A name past Excel's 31 characters is cut, so it
    cannot swell each range and cell reference of an answer. Two long names
    with one prefix stay unique, and a cut name passes a real name."""
    prefix = "Sales " + "x" * 25
    assert len(prefix) == E.SHEET_NAME_CHARS
    table = xb.sheet(_rows((1, inline("A1", "a")), (2, num("A2", 1)), (3, num("A3", 2))))
    real = prefix[:29].upper() + "~2"  # A real name that the second cut name meets
    sheets = [(prefix + "y" * 1_000_000, table), (prefix + "z" * 50, table), (real, table)]
    answer = _ok(_load(dirs, "l.xlsx", _book(sheets)))
    names = [s["name"] for s in answer["sheets"]]
    assert names == [prefix, prefix[:29] + "~3", real]
    assert len({n.casefold() for n in names}) == 3
    assert len(json.dumps(answer)) < 20_000
    ds = answer["dataset_id"]
    first = next(t for t in answer["tables"] if t["sheet"] == prefix)
    assert first["range"] == f"'{prefix}'!A2:A3"
    preview = _ok(_call(dirs, "preview", dataset_id=ds, table=first["name"]))
    assert preview["cells"] == [f"'{prefix}'!A2:A2", f"'{prefix}'!A3:A3"]
    picked = _ok(_load(dirs, "l.xlsx", _book(sheets), sheet=prefix[:29] + "~3"))
    assert [s["name"] for s in picked["sheets"]] == [prefix[:29] + "~3"]


def test_a_number_format_past_255_characters_is_no_date_and_reads_quickly(dirs: Any) -> None:
    """Security review P3-2. openpyxl's date rule costs more than linear time
    on a long crafted code. The engine skips a code past Excel's 255
    characters, and its number stays a number."""
    crafted = "[h" * 50_000
    assert E._date_kind(14, None) == "date"  # The first call imports openpyxl. Time the next.
    began = time.perf_counter()
    assert E._date_kind(164, crafted) is None
    assert time.perf_counter() - began < 0.1
    edge = "yyyy-mm-dd" + '"' + "x" * (E.FORMAT_CODE_CHARS - 12) + '"'
    assert len(edge) == E.FORMAT_CODE_CHARS and E._date_kind(164, edge) == "date"
    assert E._date_kind(164, edge + " ") is None
    styles = _STYLES.replace(b"dd/mm/yyyy hh:mm", crafted.encode())
    rows = _rows((1, inline("A1", "When")), (2, _date("A2", 45658, 2)), (3, _date("A3", 45659, 2)))
    began = time.perf_counter()
    data = _book([("D", xb.sheet(rows))], extra={"xl/styles.xml": styles})
    cols = _columns(dirs, _ok(_load(dirs, "d.xlsx", data))["dataset_id"])
    assert time.perf_counter() - began < 5
    assert cols["When"]["type"] == "integer"


# --------------------------------------------------------------------------
# Fences: done-when 7, the pins and the image.
# --------------------------------------------------------------------------


def _duckdb_imports(source: str) -> bool:
    """True when *source* imports duckdb, by statement or by name at run time."""
    for node in ast.walk(ast.parse(source)):
        if isinstance(node, ast.Import):
            names = [a.name for a in node.names]
        elif isinstance(node, ast.ImportFrom):
            names = [node.module or ""]
        elif isinstance(node, ast.Call) and getattr(
            node.func, "attr", getattr(node.func, "id", "")
        ) in ("import_module", "__import__"):
            names = [str(a.value) for a in node.args if isinstance(a, ast.Constant)]
        else:
            continue
        if any(name.split(".")[0] == "duckdb" for name in names):
            return True
    return False


def _python_files(root: Path) -> list[Path]:
    skip = {".venv", "node_modules", "__pycache__", ".git"}
    return [p for p in root.rglob("*.py") if not skip.intersection(p.relative_to(root).parts)]


def test_no_module_under_apps_or_packages_imports_duckdb_but_the_engine() -> None:
    """Done-when 7. A gateway module that imports DuckDB would take the dev package to production."""
    found = [
        p.relative_to(_REPO).as_posix()
        for root in (_REPO / "apps", _REPO / "packages")
        for p in _python_files(root)
        if _duckdb_imports(p.read_text(encoding="utf-8-sig"))
    ]
    assert found == ["apps/services/orchestrator/sandbox/data_engine.py"]


@pytest.mark.parametrize(
    ("source", "imports"),
    [
        ("import duckdb\n", True),
        ("from duckdb import connect\n", True),
        ("import os, duckdb.typing\n", True),
        ("def f():\n    import duckdb\n", True),
        ("importlib.import_module('duckdb')\n", True),
        ("__import__('duckdb')\n", True),
        ("import duckdb_extras_not\nx = 'import duckdb'\n", False),
    ],
)
def test_the_duckdb_import_check_sees(source: str, imports: bool) -> None:
    assert _duckdb_imports(source) is imports


def test_the_engine_imports_only_stdlib_duckdb_and_openpyxl() -> None:
    tree = ast.parse(_ENGINE_PATH.read_text(encoding="utf-8"))
    roots = {
        a.name.split(".")[0] for n in ast.walk(tree) if isinstance(n, ast.Import) for a in n.names
    }
    roots |= {
        n.module.split(".")[0] for n in ast.walk(tree) if isinstance(n, ast.ImportFrom) and n.module
    }
    assert roots - set(sys.stdlib_module_names) == {"duckdb", "openpyxl"}


def test_the_host_and_the_image_pin_the_same_duckdb() -> None:
    image = re.search(r"^duckdb==(\S+)", _IMAGE_LOCK.read_text(encoding="utf-8"), re.M)
    host = re.search(
        r'name = "duckdb"\nversion = "([^"]+)"', (_REPO / "uv.lock").read_text(encoding="utf-8")
    )
    dev = re.search(r'"duckdb==([^"]+)"', (_REPO / "pyproject.toml").read_text(encoding="utf-8"))
    assert image and host and dev
    assert image.group(1) == host.group(1) == dev.group(1) == duckdb.__version__


def test_the_dockerfile_copies_the_engine_read_only() -> None:
    copies = [
        args
        for word, args in _instructions(_DOCKERFILE.read_text(encoding="utf-8"))
        if word == "COPY"
    ]
    assert "--chmod=0444 sandbox/data_engine.py /opt/sandbox/data_engine.py" in copies


# --------------------------------------------------------------------------
# The Docker half. Deselected by default (pyproject.toml); the workflow
# .github/workflows/sandbox-docker.yml runs it and fails on any skip.
# --------------------------------------------------------------------------

#: The helper of both Docker scripts. ``call`` runs one verb through a small
#: wrapper, so ``ru_maxrss`` of the wrapper's children is the peak of that verb:
#: the engine and its forked child.
_DOCKER_HELPER = r"""
import importlib.util, json, os, resource, subprocess, sys, time
from pathlib import Path
run = Path("/workspace/.run")
(run / "requests").mkdir(exist_ok=True)
(run / "src").mkdir(exist_ok=True)
WRAP = "import resource, subprocess, sys; r = subprocess.run(sys.argv[1:]); print(r.returncode, resource.getrusage(resource.RUSAGE_CHILDREN).ru_maxrss)"
def call(verb, **req):
    rid = f"r{len(list((run / 'requests').iterdir()))}"
    (run / "requests" / f"{rid}.json").write_text(json.dumps(req))
    started = time.monotonic()
    done = subprocess.run(["python3", "-I", "-c", WRAP, "python3", "-I", "/opt/sandbox/data_engine.py", verb, rid], capture_output=True, text=True)
    code, rss = done.stdout.split()
    path = run / "answers" / f"{rid}.json"
    answer = json.loads(path.read_text()) if path.exists() else {}
    answer["_rc"], answer["_peak_mb"] = int(code), int(rss) // 1024
    answer["_s"] = round(time.monotonic() - started, 1)
    return answer
def cgroup_peak_mb():
    for name in ("/sys/fs/cgroup/memory.peak", "/sys/fs/cgroup/memory/memory.max_usage_in_bytes"):
        try:
            return int(Path(name).read_text()) // 2**20
        except OSError:
            pass
    return None
MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
def big_book(path, rows, bomb_mb=0):
    # A workbook of raw XML, streamed: id, amount, flag, a shared string and a date.
    # With bomb_mb, the sheet is that many MB of one inline string instead.
    import zipfile
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("_rels/.rels", f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="{REL}/officeDocument" Target="xl/workbook.xml"/></Relationships>')
        zf.writestr("xl/workbook.xml", f'<workbook xmlns="{MAIN}" xmlns:r="{REL}"><sheets><sheet name="Data" sheetId="1" r:id="rId1"/></sheets></workbook>')
        zf.writestr("xl/_rels/workbook.xml.rels", f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="{REL}/worksheet" Target="worksheets/sheet1.xml"/><Relationship Id="rId2" Type="{REL}/sharedStrings" Target="sharedStrings.xml"/></Relationships>')
        zf.writestr("xl/styles.xml", f'<styleSheet xmlns="{MAIN}"><cellXfs count="2"><xf numFmtId="0"/><xf numFmtId="14"/></cellXfs></styleSheet>')
        zf.writestr("xl/sharedStrings.xml", f'<sst xmlns="{MAIN}">' + "".join(f"<si><t>C{i}</t></si>" for i in range(50)) + "</sst>")
        with zf.open("xl/worksheets/sheet1.xml", "w", force_zip64=True) as fh:
            fh.write(f'<worksheet xmlns="{MAIN}"><sheetData>'.encode())
            if bomb_mb:
                fh.write(b'<row r="1"><c r="A1" t="inlineStr"><is><t>')
                for _ in range(bomb_mb):
                    fh.write(b"A" * 2**20)
                fh.write(b"</t></is></c></row>")
            else:
                fh.write(b'<row r="1"><c r="A1" t="inlineStr"><is><t>id</t></is></c><c r="B1" t="inlineStr"><is><t>amount</t></is></c><c r="C1" t="inlineStr"><is><t>flag</t></is></c><c r="D1" t="inlineStr"><is><t>who</t></is></c><c r="E1" t="inlineStr"><is><t>when</t></is></c></row>')
                for start in range(2, rows + 2, 5000):
                    fh.write("".join(
                        f'<row r="{i}"><c r="A{i}"><v>{i}</v></c><c r="B{i}"><v>{i % 997}.5</v></c><c r="C{i}" t="b"><v>{i % 2}</v></c><c r="D{i}" t="s"><v>{i % 50}</v></c><c r="E{i}" s="1"><v>{45000 + i % 365}</v></c></row>'
                        for i in range(start, min(start + 5000, rows + 2))).encode())
            fh.write(b"</sheetData></worksheet>")
"""

_DOCKER_SCRIPT = r"""
set -e
echo "uid: $(id -u)"
echo "engine: $(stat -c '%u %a' /opt/sandbox/data_engine.py)"
cd /tmp
python3 -I - <<'PY'
HELPER
(run / "src" / "Invoices.csv").write_text(INVOICES_TEXT, encoding="utf-8")
(run / "other.csv").write_text("a\n1\n", encoding="utf-8")
out = {}
load = call("load", source="src/Invoices.csv")
ds = load["dataset_id"]
out["load"] = load["ok"]
out["query"] = call("query", dataset_id=ds, sql='SELECT sum("Amount") FROM invoices')["rows"]
out["preview"] = call("preview", dataset_id=ds, table="invoices", limit=1)["cells"]
out["profile"] = call("profile", dataset_id=ds, table="invoices", column="Amount")["ok"]
out["export"] = call("export", dataset_id=ds, sql="SELECT 1 AS a", format="xlsx")["file"]
from openpyxl import Workbook
book = Workbook()
sheet = book.active
sheet.title = "Q1 Sales"
for line in (["Sales report"], [], ["Region", "Amount"], ["North", 10], ["South", 30], ["Total", 40]):
    sheet.append(line)
book.save(run / "src" / "r.xlsx")
loaded = call("load", source="src/r.xlsx")
out["xlsx"] = [t["name"] for t in loaded.get("tables", [])] or loaded
out["xlsx sum"] = call("query", dataset_id=loaded.get("dataset_id"), sql="SELECT sum(Amount) FROM q1_sales").get("rows")
out["two"] = call("query", dataset_id=ds, sql="SELECT 1; SELECT 2")["error"]
spec = importlib.util.spec_from_file_location("data_engine", "/opt/sandbox/data_engine.py")
engine = importlib.util.module_from_spec(spec)
sys.modules["data_engine"] = engine
spec.loader.exec_module(engine)
d = f"/workspace/.data/{ds}"
pq = f"{d}/invoices.parquet"
copy = f"COPY (SELECT 99 AS a) TO '{pq}'"
cases = {
    "read_csv": "SELECT * FROM read_csv('/workspace/.run/other.csv')",
    "read_text": "SELECT * FROM read_text('/etc/passwd')",
    "glob": f"SELECT * FROM glob('{d}/*')",
    "getenv": "SELECT getenv('PATH')",
    "attach": f"ATTACH '{d}/x.db'",
    "copy into the dataset dir": f"COPY (SELECT 1 AS a) TO '{d}/w.csv'",
    "install": "INSTALL httpfs",
    "load": "LOAD httpfs",
    "set": "SET threads = 4",
    "pragma": "PRAGMA threads = 4",
    "overwrite parquet": f"{copy} (FORMAT parquet)",
    "overwrite parquet, no tmp file": f"{copy} (FORMAT parquet, USE_TMP_FILE false)",
    "overwrite parquet, tmp file": f"{copy} (FORMAT parquet, USE_TMP_FILE true)",
    "overwrite csv, no tmp file": f"{copy} (FORMAT csv, USE_TMP_FILE false)",
    "overwrite csv": f"{copy} (FORMAT csv)",
    "overwrite true": f"{copy} (FORMAT parquet, OVERWRITE true)",
    "overwrite or ignore": f"{copy} (FORMAT parquet, OVERWRITE_OR_IGNORE true)",
}
before = {p.name: p.read_bytes() for p in Path(d).iterdir()}
allowed = []
for label, sql in cases.items():
    con, _ = engine._open_locked(engine.Dirs(), ds)
    try:
        con.execute(sql).fetchall()
        allowed.append(label)
    except Exception:
        pass
    assert con.execute('SELECT count(*), sum("Amount") FROM invoices').fetchone()[0] == 4, label
out["lock_allowed"] = allowed
out["dataset_same"] = {p.name: p.read_bytes() for p in Path(d).iterdir()} == before
out["modes"] = sorted(f"{p.name} {oct(p.stat().st_mode & 0o777)}" for p in [Path(d), *Path(d).iterdir()])
print("RESULT " + json.dumps(out))
PY
"""


def _docker_run(image: str, script: str, *extra: str, size: str = "64m") -> Any:
    writable = []
    for target in ("/workspace/.run", "/workspace/.data", "/workspace/outputs"):
        writable += ["--tmpfs", f"{target}:rw,nosuid,nodev,size={size},mode=1777"]
    return _docker(
        "run", *_RUN_FLAGS, *writable, *extra, "--user", "1000:1000", image,
        "bash", "-o", "pipefail", "-c", script, timeout=900,
    )  # fmt: skip


def _result(result: Any) -> dict[str, Any]:
    assert result.returncode == 0, f"stdout:\n{result.stdout}\nstderr:\n{result.stderr}"
    line = next(line for line in result.stdout.splitlines() if line.startswith("RESULT "))
    return dict(json.loads(line[7:]))


@pytest.mark.sandbox_docker
def test_each_verb_and_each_lock_case_runs_in_the_image_as_uid_1000(
    coding_sandbox_image: str,  # noqa: F811 — the fixture, imported above
) -> None:
    """Done-when 5 and 6, and rule 4 in the image: uid 1000, read-only, no network."""
    script = _DOCKER_SCRIPT.replace("HELPER", _DOCKER_HELPER).replace(
        "INVOICES_TEXT", repr(INVOICES)
    )
    result = _docker_run(coding_sandbox_image, script)
    lines = result.stdout.splitlines()
    assert "uid: 1000" in lines and "engine: 0 444" in lines, result.stdout
    out = _result(result)
    assert out["load"] is True and out["query"] == [[AMOUNT_SUM]]
    assert out["preview"] == ["invoices!A3:G3"] and out["profile"] is True
    assert out["export"] == "outputs/export.xlsx"
    # WS-43y1b: a workbook that openpyxl wrote, with a title row and a total row.
    assert out["xlsx"] == ["q1_sales", "q1_sales__totals"], out["xlsx"]
    assert out["xlsx sum"] == [[40]] and out["two"] == "one_select"
    assert out["lock_allowed"] == [], out
    assert out["dataset_same"] is True
    ds = next(m.split()[0] for m in out["modes"] if m.startswith("ds_"))
    assert out["modes"] == sorted([f"{ds} 0o555", "invoices.parquet 0o444", "manifest.json 0o444"])


_MEMORY_SCRIPT = r"""
set -e
cd /tmp
python3 -I - <<'PY'
HELPER
out = {}
with open(run / "src" / "mixed.csv", "w") as fh:
    fh.write("id,amount,flag,who,when\n")
    for i in range(1, 1_000_001):
        when = f"2024-{i % 12 + 1:02d}-{i % 28 + 1:02d}" if i % 4 == 0 else ""
        fh.write(f"{i},{i % 997}.5,{'yes' if i % 2 else 'no'},C{i % 50},{when}\n")
with open(run / "src" / "ones.csv", "w") as fh:
    fh.write("1\n" * (25 * 2**20 // 2))
out["bytes"] = {p.name: p.stat().st_size for p in (run / "src").iterdir()}
ds = None
for name in ("mixed.csv", "ones.csv"):
    a = call("load", source=f"src/{name}", timeout_seconds=300)
    out[name] = {k: a.get(k) for k in ("ok", "error", "message", "_rc", "_peak_mb", "_s")}
    out[name]["rows"] = a["tables"][0]["rows"] if a.get("ok") else None
    out[name]["types"] = [c["type"] for c in a["tables"][0]["columns"]] if a.get("ok") else None
    if a.get("ok"):
        ds = a["dataset_id"]
(run / "src" / "t.csv").write_text("a,b\n1,x\n2,y\n")
small = call("load", source="src/t.csv")["dataset_id"]
attacks = {
    "len(range 1e8)": "SELECT len(range(0, 100000000)) AS n",
    "range 4e8": "SELECT range(0, 400000000) AS r",
    "repeat 1.5e9": "SELECT repeat('x', 1500000000) AS s",
    "string_agg 1.5e8": "SELECT length(string_agg(range::VARCHAR, ',')) FROM range(150000000)",
}
for label, sql in attacks.items():
    a = call("query", dataset_id=small, sql=sql, timeout_seconds=20)
    out[label] = {k: a.get(k) for k in ("ok", "error", "_rc", "_peak_mb", "_s")}
for label, (width, name_len, rows) in {"wide 1000x2000": (1000, 8, 2000), "names 400x1000": (400, 1000, 10)}.items():
    head = ",".join(f"h{i}_" + "x" * name_len for i in range(width))
    body = "".join(",".join(str(r * j % 97) + (".5" if j % 3 == 0 else "") for j in range(width)) + "\n" for r in range(rows))
    (run / "src" / f"{label[:5]}.csv").write_text(head + "\n" + body)
    a = call("load", source=f"src/{label[:5]}.csv", timeout_seconds=300)
    out[label] = {k: a.get(k) for k in ("ok", "error", "_rc", "_peak_mb", "_s")}
def scratch():
    return sorted(p.name for p in Path("/tmp").iterdir())
before = scratch()
amp = "SELECT repeat('&', 2000) AS s FROM range(200000)"
out["xlsx amp"] = {k: v for k, v in call("export", dataset_id=small, sql=amp, format="xlsx", timeout_seconds=120).items() if k in ("ok", "error", "_rc")}
slow = "SELECT range, repeat('y', 50) AS s FROM range(5000000)"
out["xlsx timer"] = {k: v for k, v in call("export", dataset_id=small, sql=slow, format="xlsx", timeout_seconds=5).items() if k in ("ok", "error", "_rc")}
out["tmp after xlsx"] = [n for n in scratch() if n not in before]
out["outputs after xlsx"] = sorted(p.name for p in Path("/workspace/outputs").iterdir())
out["next verb"] = call("query", dataset_id=small, sql="SELECT 1 AS one").get("rows")
def stage_peak(label, name, text):
    (run / "src" / name).write_text(text)
    peak, done = [0], [False]
    def watch():
        while not done[0]:
            size = sum(f.stat().st_size for f in Path("/workspace/.data").rglob("*") if f.is_file())
            peak[0] = max(peak[0], size)
            time.sleep(0.1)
    import threading
    th = threading.Thread(target=watch)
    th.start()
    a = call("load", source=f"src/{name}", timeout_seconds=300)
    done[0] = True
    th.join()
    out[label] = {k: a.get(k) for k in ("ok", "error", "_rc", "_s")}
    out[label]["stage_peak_mb"] = peak[0] // 2**20
    out[label]["src_mb"] = round(len(text) / 2**20, 1)
pad_header = ",".join(f"c{i}" for i in range(1000)) + "\n"
stage_peak("pad wide header", "pw.csv", pad_header + "1\n" * 2_000_000)
stage_peak("pad late wide row", "pl.csv", "a\n" + "1\n" * 1_999_998 + ",".join(["9"] * 1000) + "\n")
out["mixed sum"] = ds and call("query", dataset_id=ds, sql="SELECT count(*), sum(amount) FROM mixed", timeout_seconds=60).get("rows")
# WS-43y1b: a large workbook streams, and a zip bomb stops before its part unpacks.
for p in (run / "src").iterdir():
    p.unlink()
big_book(run / "src" / "big.xlsx", 800_000)
out["xlsx bytes"] = (run / "src" / "big.xlsx").stat().st_size
a = call("load", source="src/big.xlsx", timeout_seconds=300)
out["big.xlsx"] = {k: a.get(k) for k in ("ok", "error", "message", "_rc", "_peak_mb", "_s")}
out["big.xlsx"]["rows"] = a["tables"][0]["rows"] if a.get("ok") else None
out["big.xlsx"]["types"] = [c["type"] for c in a["tables"][0]["columns"]] if a.get("ok") else None
out["big.xlsx sum"] = a.get("ok") and call("query", dataset_id=a["dataset_id"], sql="SELECT count(*), sum(amount) FROM data", timeout_seconds=60).get("rows")
big_book(run / "src" / "bomb.xlsx", 0, bomb_mb=201)
out["bomb bytes"] = (run / "src" / "bomb.xlsx").stat().st_size
out["bomb"] = {k: v for k, v in call("load", source="src/bomb.xlsx").items() if k in ("error", "_rc", "_peak_mb", "_s")}
out["cgroup_peak_mb"] = cgroup_peak_mb()
print("RESULT " + json.dumps(out))
PY
"""


@pytest.mark.sandbox_docker
def test_the_engine_answers_under_a_1_gib_container(
    coding_sandbox_image: str,  # noqa: F811 — the fixture, imported above
) -> None:
    """P1-D and P2-E, in the image with ``--memory 1g`` as the broker runs it (§7.1 rule 6).

    A 24 MB, 1M-row CSV loads. A 25 MB file of ``1`` lines answers (it passes
    the row cap). Each query that eats memory answers with an error inside
    its time cap plus the kill margin, and the engine never dies silently.
    """
    script = _MEMORY_SCRIPT.replace("HELPER", _DOCKER_HELPER)
    extra = ("--memory", "1g", "--memory-swap", "1g")
    out = _result(_docker_run(coding_sandbox_image, script, *extra, size="160m"))
    print(json.dumps(out, indent=1))  # The record of the peaks (-s shows it).
    mixed = out["mixed.csv"]
    assert mixed["ok"] and mixed["rows"] == 1_000_000 and mixed["_rc"] == 0, mixed
    assert mixed["types"] == ["integer", "decimal", "boolean", "text", "date"]
    assert out["mixed sum"] == [[1_000_000, out["mixed sum"][0][1]]]
    ones = out["ones.csv"]
    assert ones["_rc"] == 0 and ones["error"] == "too_many_rows", ones
    for label in ("len(range 1e8)", "range 4e8", "repeat 1.5e9", "string_agg 1.5e8"):
        attack = out[label]
        assert attack["_rc"] == 0 and attack["ok"] is False, (label, attack)
        assert attack["error"] in ("memory", "time", "sql_error"), (label, attack)
        assert attack["_s"] <= 20 + E.KILL_MARGIN_SECONDS + 5, (label, attack)
        # The parent's watch, or the kernel at the cgroup, stops the child.
        assert attack["_peak_mb"] <= 1024, (label, attack)
    for label in ("wide 1000x2000", "names 400x1000"):
        assert out[label]["ok"] is True and out[label]["_rc"] == 0, (label, out[label])
        # H-290: half the watch. With one database for the whole load, the
        # names shape took 388 to 507 MB locally and 771 MB on a CI runner.
        # With a database for each statement, names takes 186 to 204 MB and
        # wide takes 283 to 298 MB.
        assert out[label]["_peak_mb"] < 384, (label, out[label])
    assert out["xlsx amp"]["error"] == "too_large" and out["xlsx timer"]["error"] == "time"
    assert out["tmp after xlsx"] == [] and out["outputs after xlsx"] == []
    assert out["next verb"] == [[1]]
    # Round 3: a file too wide for its size stops at the stage budget.
    for label in ("pad wide header", "pad late wide row"):
        pad = out[label]
        assert pad["error"] == "bad_file" and pad["_rc"] == 0, (label, pad)
        assert pad["stage_peak_mb"] < 200 and pad["_s"] < 60, (label, pad)
    # WS-43y1b: 800,000 rows of a workbook load inside the load's time cap and
    # the container. A zip bomb answers at once, and unpacks nothing.
    book = out["big.xlsx"]
    assert book["ok"] and book["rows"] == 800_000 and book["_rc"] == 0, book
    assert book["types"] == ["integer", "decimal", "boolean", "text", "date"], book
    assert book["_peak_mb"] < 768 and book["_s"] < E.LOAD_TIMEOUT_SECONDS, book
    assert out["big.xlsx sum"][0][0] == 800_000
    bomb = out["bomb"]
    assert bomb["error"] == "too_large" and bomb["_s"] < 10 and bomb["_peak_mb"] < 200, bomb


_GROW_SCRIPT = r"""
set -e
cd /tmp
python3 -I - <<'PY'
HELPER
(run / "src" / "t.csv").write_text("a,b\n1,x\n")
(run / "requests" / "r0.json").write_text(json.dumps({"source": "src/t.csv"}))
os.environ.pop("DATA_ENGINE_FAULT")
subprocess.run(["python3", "-I", "/opt/sandbox/data_engine.py", "load", "r0"], check=True)
ds = json.loads((run / "answers" / "r0.json").read_text())["dataset_id"]
os.environ["DATA_ENGINE_FAULT"] = "grow"
a = call("query", dataset_id=ds, sql="SELECT 1", timeout_seconds=60)
print("RESULT " + json.dumps({k: a.get(k) for k in ("ok", "error", "_rc", "_peak_mb", "_s")}))
PY
"""


@pytest.mark.sandbox_docker
def test_the_watch_kills_a_child_that_grows(
    coding_sandbox_image: str,  # noqa: F811 — the fixture, imported above
) -> None:
    """Round 2: the parent kills a child at 75% of the container's memory.

    The fault hook makes the child take 32 MB each 50 ms. With no watch, the
    kernel kills it at the cgroup, near 1 GiB, so the peak catches a missing
    watch.
    """
    script = _GROW_SCRIPT.replace("HELPER", _DOCKER_HELPER)
    extra = ("--memory", "1g", "--memory-swap", "1g", "-e", "DATA_ENGINE_FAULT=grow")
    out = _result(_docker_run(coding_sandbox_image, script, *extra))
    print(json.dumps(out))
    assert out["ok"] is False and out["error"] == "memory" and out["_rc"] == 0, out
    assert out["_peak_mb"] < 850, out
