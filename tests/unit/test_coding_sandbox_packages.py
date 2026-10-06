"""WS43-F25 — the coding sandbox holds every package that its code imports.

Spec: ``project-docs/specs/maf_coding_engine.md`` §7.2 (the image and the
rule for a new package) and §10 (the fence row WS43-F25).

The container has no network at run time (§7.1 rule 6). So a package that is
not in the image is a package that no script can use: ``pip install`` fails,
and the agent's rule 6 forbids it anyway. This file fails when the code that
runs in the sandbox needs a package that the image does not hold.

**The static half** runs in the default unit job. It fails when:

- ``sandbox/requirements.in`` (the list) and :data:`_SANDBOX_MODULES` (the
  import name of each package) name different packages,
- ``sandbox/requirements.txt`` (the lock) does not pin a name of the list, or
  pins as a direct requirement a name that the list dropped (a stale lock),
- the lock pins a browser engine, a GPU or ML-scale library, or a package
  that needs the network at run time (:data:`_LEFT_OUT`),
- a library that the sandbox section of the instructions names
  (``acb_skills.addendum.SANDBOX_LIBRARIES``) is not in the list,
- a Python import in the sandbox code is not the standard library, a module
  of the list, or a file of the same skill. The sandbox code is:
  the known-good scripts of ``evals/coding_engine/scripted.py``, and the
  scripts and ``python`` blocks of each prebuilt agent skill
  (``skills/<domain>/<skill>/`` and ``apps/agents/*/skills/``).

``skills/upstream/`` is out of scope on purpose. It is a review mirror, and
nothing loads it (``skills/upstream/README.md``). Its ``pdf``, ``docx``,
``pptx`` and ``xlsx`` skills ask for LibreOffice, Poppler, Tesseract and npm
packages too. A skill that is adopted into ``skills/<domain>/`` comes under
this fence.

Each checker also runs on synthetic bad input, so a checker that goes blind
is a red test.

**The Docker half** carries the ``sandbox_docker`` marker. It imports every
module of the list in the image, as uid 1000, read-only and with no network,
and makes one PNG, one ``.docx``, one ``.pptx``, one ``.xlsx`` and one PDF.
``.github/workflows/sandbox-docker.yml`` runs it and fails on any skip.
"""

from __future__ import annotations

import ast
import re
import sys
from collections.abc import Iterable
from pathlib import Path

import pytest

from tests.unit.test_coding_sandbox_image import (  # noqa: F401 — a fixture, used by name
    _ORCH,
    _pinned_names,
    _run_in_image,
    coding_sandbox_image,
)

_REPO = Path(__file__).resolve().parents[2]
_LIST = _ORCH / "sandbox" / "requirements.in"
_LOCK = _ORCH / "sandbox" / "requirements.txt"
_LIST_REL = "apps/services/orchestrator/sandbox/requirements.in"

#: Each package of ``requirements.in``, by its normalised distribution name,
#: and the module that it imports as. A new package needs a reviewed PR that
#: adds it here AND to ``requirements.in`` (§7.2).
_SANDBOX_MODULES: dict[str, str] = {
    # The first set of §7.2.
    "pandas": "pandas",
    "numpy": "numpy",
    "openpyxl": "openpyxl",
    "matplotlib": "matplotlib",
    "pypdf": "pypdf",
    "fpdf2": "fpdf",
    "markdown": "markdown",
    "tabulate": "tabulate",
    "pyyaml": "yaml",
    "python-dateutil": "dateutil",
    "requests": "requests",
    # The second set (2026-10-05): business documents.
    "python-docx": "docx",
    "python-pptx": "pptx",
    "xlsxwriter": "xlsxwriter",
    "reportlab": "reportlab",
    "pdfplumber": "pdfplumber",
    "pillow": "PIL",
    # Analysis and charts.
    "scipy": "scipy",
    "seaborn": "seaborn",
    "networkx": "networkx",
    # HTML, XML and templates.
    "beautifulsoup4": "bs4",
    "lxml": "lxml",
    "jinja2": "jinja2",
}

#: Packages that the image must not pin, directly or as a dependency, and why.
#: To take one in is a decision for a reviewed PR that edits this table.
_LEFT_OUT: dict[str, str] = {
    "playwright": "a browser engine",
    "pyppeteer": "a browser engine",
    "selenium": "drives a browser engine",
    "kaleido": "plotly's static export: v1 drives a Chrome, v0.2 bundles a Chromium",
    "torch": "GPU and ML scale",
    "tensorflow": "GPU and ML scale",
    "jax": "GPU and ML scale",
    "scikit-learn": "ML scale, and no agent uses it",
    "onnxruntime": "an ML runtime (markitdown pulls it in through magika)",
    "pytesseract": "needs the Tesseract binary and its language data",
    "pdf2image": "needs the Poppler binaries; pypdfium2 renders a page instead",
}

#: The roots of the prebuilt agent skills. ``skills/upstream`` is a mirror that
#: nothing loads (see the module docstring).
_SKILL_ROOTS = (_REPO / "skills", _REPO / "apps" / "agents")
_NOT_A_SKILL_ROOT = (_REPO / "skills" / "upstream",)

_FENCE_RE = re.compile(r"^```(?:python|py|python3)[ \t]*\n(.*?)^```", re.M | re.S)
_IMPORT_LINE_RE = re.compile(r"^\s*(?:from\s+([A-Za-z_]\w*)|import\s+([A-Za-z_]\w*))", re.M)


def _norm(name: str) -> str:
    return re.sub(r"[-_.]+", "-", name).lower()


# --------------------------------------------------------------------------
# Parsers and checkers. Each returns a list of problems, empty when good.
# --------------------------------------------------------------------------


def _listed_names(text: str) -> set[str]:
    """The names of ``requirements.in``: one bare name a line, and comments."""
    names = set()
    for raw in text.splitlines():
        line = raw.split("#", 1)[0].strip()
        if line:
            names.add(_norm(re.split(r"[\s<>=!~;\[]", line, maxsplit=1)[0]))
    return names


def _list_problems(text: str) -> list[str]:
    """Each line of the list is one bare name, and the list matches the table."""
    problems: list[str] = []
    for raw in text.splitlines():
        line = raw.split("#", 1)[0].strip()
        if line and not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]*", line):
            problems.append(f"not a bare name (the lock holds the pin): {line!r}")
    names = _listed_names(text)
    if missing := sorted(set(_SANDBOX_MODULES) - names):
        problems.append(f"_SANDBOX_MODULES names these and requirements.in does not: {missing}")
    if extra := sorted(names - set(_SANDBOX_MODULES)):
        problems.append(f"requirements.in names these and _SANDBOX_MODULES does not: {extra}")
    return problems


def _direct_names(lock: str) -> set[str]:
    """The names that the lock says came from the list (``# via -r …in``)."""
    direct = set()
    blocks = re.split(r"\n(?=[A-Za-z0-9])", lock)
    for block in blocks:
        head = block.split(None, 1)[0] if block.strip() else ""
        if "==" in head and f"-r {_LIST_REL}" in block:
            direct.add(_norm(head.split("==", 1)[0]))
    return direct


def _lock_problems(listed: set[str], lock: str) -> list[str]:
    """The lock pins every listed name, and nothing that the list dropped."""
    problems: list[str] = []
    head = lock.splitlines()[:2]
    if len(head) < 2 or _LIST_REL not in head[1]:
        problems.append(f"the lock header does not name {_LIST_REL}: recompile it")
    pinned = _pinned_names(lock)
    if missing := sorted(listed - pinned):
        problems.append(f"the lock does not pin these listed names: {missing}")
    if stale := sorted(_direct_names(lock) - listed):
        problems.append(f"the lock pins these as direct, and the list dropped them: {stale}")
    if unmarked := sorted(listed - _direct_names(lock)):
        problems.append(f"the lock does not mark these as from the list: {unmarked}")
    return problems


def _left_out_problems(lock: str) -> list[str]:
    pinned = _pinned_names(lock)
    return [f"the lock pins {name}: {why}" for name, why in _LEFT_OUT.items() if name in pinned]


def _top_imports(source: str) -> set[str]:
    """The top-level modules that *source* imports. A relative import is local."""
    try:
        tree = ast.parse(source)
    except SyntaxError:
        # A block of a SKILL.md can be a sketch. Read its import lines.
        return {a or b for a, b in _IMPORT_LINE_RE.findall(source)}
    out: set[str] = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            out.update(alias.name.split(".", 1)[0] for alias in node.names)
        elif isinstance(node, ast.ImportFrom) and node.level == 0 and node.module:
            out.add(node.module.split(".", 1)[0])
    return out


def _allowed_modules() -> set[str]:
    return set(sys.stdlib_module_names) | set(_SANDBOX_MODULES.values())


def _import_problems(label: str, source: str, local: Iterable[str] = ()) -> list[str]:
    missing = sorted(_top_imports(source) - _allowed_modules() - set(local))
    return [f"{label} imports {m}, and the image does not hold it" for m in missing]


def _skill_dirs() -> list[Path]:
    """Each prebuilt agent skill: a folder with a ``SKILL.md``."""
    found: list[Path] = []
    for root in _SKILL_ROOTS:
        for skill_md in sorted(root.rglob("SKILL.md")):
            folder = skill_md.parent
            if any(folder.is_relative_to(n) for n in _NOT_A_SKILL_ROOT):
                continue
            if root.name == "agents" and "skills" not in folder.relative_to(root).parts:
                continue
            found.append(folder)
    return found


def _skill_problems(folder: Path, repo: Path = _REPO) -> list[str]:
    """The scripts and the ``python`` blocks of one skill import only the image.

    A skill's own ``tests/`` run on the host under pytest, so they are not
    sandbox code.
    """
    files = [p for p in folder.rglob("*.py") if "tests" not in p.relative_to(folder).parts]
    local = {p.stem for p in files} | {p.parent.name for p in files}
    rel = folder.relative_to(repo).as_posix()
    problems: list[str] = []
    for path in files:
        source = path.read_text(encoding="utf-8")
        problems += _import_problems(f"{rel}/{path.relative_to(folder).as_posix()}", source, local)
    text = (folder / "SKILL.md").read_text(encoding="utf-8")
    for i, block in enumerate(_FENCE_RE.findall(text), 1):
        problems += _import_problems(f"{rel}/SKILL.md block {i}", block, local)
    return problems


def _eval_scripts() -> dict[str, str]:
    """The scripts that the known-good sequences run in the sandbox."""
    from evals.coding_engine import scripted

    scripts = {
        name: value
        for name, value in vars(scripted).items()
        if name.endswith("_PY") and isinstance(value, str)
    }
    for i, block in enumerate(_FENCE_RE.findall(scripted.SKILL_MD), 1):
        scripts[f"SKILL_MD block {i}"] = block
    return scripts


# --------------------------------------------------------------------------
# The static half — the real files.
# --------------------------------------------------------------------------


def _list_text() -> str:
    return _LIST.read_text(encoding="utf-8")


def _lock_text() -> str:
    return _LOCK.read_text(encoding="utf-8")


def test_the_list_and_the_table_name_the_same_packages() -> None:
    assert _list_problems(_list_text()) == []


def test_the_lock_is_compiled_from_the_list() -> None:
    assert _lock_problems(_listed_names(_list_text()), _lock_text()) == []


def test_the_lock_pins_no_browser_no_gpu_and_no_ml_runtime() -> None:
    assert _left_out_problems(_lock_text()) == []


def test_the_sandbox_section_names_only_packages_of_the_image() -> None:
    from acb_skills import addendum

    known = set(_SANDBOX_MODULES) | set(_SANDBOX_MODULES.values())
    named = list(addendum.SANDBOX_LIBRARIES)
    assert named, "SANDBOX_LIBRARIES is empty"
    assert [n for n in named if n not in known] == []
    section = addendum.render_run_sections({"run_command"})
    assert all(n in section for n in named), section[:400]


def test_the_eval_scripts_import_only_the_image() -> None:
    scripts = _eval_scripts()
    assert {"CHART_PY", "EXPORT_PY", "BURNDOWN_PY"} <= set(scripts), sorted(scripts)
    problems = [p for label, src in scripts.items() for p in _import_problems(label, src)]
    assert problems == []


def test_every_prebuilt_skill_imports_only_the_image() -> None:
    skills = _skill_dirs()
    assert skills, "no prebuilt skill found: the walk went blind"
    assert [p for folder in skills for p in _skill_problems(folder)] == []


# --------------------------------------------------------------------------
# The checkers still see: one synthetic bad input for each clause.
# --------------------------------------------------------------------------


def _good_list() -> str:
    return "# a comment\n" + "\n".join(_SANDBOX_MODULES) + "\n"


def test_the_list_check_accepts_the_table() -> None:
    assert _list_problems(_good_list()) == []


@pytest.mark.parametrize(
    "text",
    [
        pytest.param(_good_list().replace("seaborn\n", ""), id="a-name-missing"),
        pytest.param(_good_list() + "plotly\n", id="a-name-not-in-the-table"),
        pytest.param(_good_list().replace("pandas\n", "pandas==3.0.6\n"), id="a-pin-in-the-list"),
    ],
)
def test_the_list_check_refuses(text: str) -> None:
    assert _list_problems(text)


def _lock(*entries: tuple[str, str]) -> str:
    head = f"# This file was autogenerated by uv via the following command:\n#    uv pip compile {_LIST_REL} --generate-hashes\n"
    body = "".join(
        f"{pin} \\\n    --hash=sha256:{'a' * 64}\n    # via {via}\n" for pin, via in entries
    )
    return head + body


def test_the_lock_check_accepts_a_fresh_lock() -> None:
    lock = _lock(("pandas==3.0.6", f"-r {_LIST_REL}"), ("six==1.17.0", "python-dateutil"))
    assert _lock_problems({"pandas"}, lock) == []


@pytest.mark.parametrize(
    ("listed", "lock"),
    [
        pytest.param(
            {"pandas", "seaborn"}, _lock(("pandas==3.0.6", f"-r {_LIST_REL}")), id="not-pinned"
        ),
        pytest.param(
            {"pandas"},
            _lock(("pandas==3.0.6", f"-r {_LIST_REL}"), ("plotly==7.1.0", f"-r {_LIST_REL}")),
            id="stale-direct-pin",
        ),
        pytest.param(
            {"pandas"}, _lock(("pandas==3.0.6", "seaborn")), id="pinned-but-not-from-the-list"
        ),
        pytest.param(
            {"pandas"},
            _lock(("pandas==3.0.6", f"-r {_LIST_REL}")).replace(_LIST_REL, "-", 1),
            id="compiled-from-stdin",
        ),
    ],
)
def test_the_lock_check_refuses(listed: set[str], lock: str) -> None:
    assert _lock_problems(listed, lock)


def test_the_left_out_check_refuses_a_transitive_ml_runtime() -> None:
    lock = _lock(("markitdown==0.1.3", f"-r {_LIST_REL}"), ("onnxruntime==1.23.0", "magika"))
    assert _left_out_problems(lock)


@pytest.mark.parametrize(
    "source",
    [
        pytest.param("import plotly.express as px\n", id="import"),
        pytest.param("from sklearn.linear_model import LinearRegression\n", id="from-import"),
        pytest.param("import json, torch\n", id="second-name"),
        pytest.param("def f():\n    import kaleido\n", id="nested"),
        pytest.param("import six\n", id="a-dependency-not-a-listed-name"),
        pytest.param("this is not python\nimport plotly\n", id="a-sketch"),
    ],
)
def test_the_import_check_refuses(source: str) -> None:
    assert _import_problems("x", source)


def test_the_import_check_accepts_stdlib_listed_and_local() -> None:
    source = (
        "import json\nfrom pathlib import Path\nimport docx\nfrom PIL import Image\n"
        "from . import sibling\nimport helpers\n"
    )
    assert _import_problems("x", source, local={"helpers"}) == []


def test_the_skill_check_sees_a_script_and_a_block(tmp_path: Path) -> None:
    skill = tmp_path / "skill"
    (skill / "scripts").mkdir(parents=True)
    (skill / "scripts" / "chart.py").write_text("import plotly\n", encoding="utf-8")
    (skill / "SKILL.md").write_text("# s\n\n```python\nimport torch\n```\n", encoding="utf-8")
    (skill / "tests").mkdir()
    (skill / "tests" / "test_chart.py").write_text("import pytest\n", encoding="utf-8")
    problems = _skill_problems(skill, repo=tmp_path)
    assert any("plotly" in p for p in problems), problems
    assert any("torch" in p for p in problems), problems
    assert not any("pytest" in p for p in problems), problems


# --------------------------------------------------------------------------
# The Docker half. Deselected by default (pyproject.toml); the workflow
# .github/workflows/sandbox-docker.yml runs it and fails on any skip.
# --------------------------------------------------------------------------

_SMOKE_SCRIPT = """\
set -e
cd /tmp
python - <<'PY'
import importlib
for name in {modules!r}:
    module = importlib.import_module(name)
    assert module.__file__, f"{{name}} imported as an empty namespace package"
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
import seaborn as sns
frame = pd.DataFrame({{"who": ["A", "B", "C"], "open": [3, 5, 2]}})
sns.barplot(data=frame, x="who", y="open")
plt.savefig("chart.png")
import docx
doc = docx.Document()
doc.add_heading("Status", 0)
doc.add_paragraph("\\u20b9500 per user")
doc.add_picture("chart.png")
doc.save("r.docx")
import pptx
deck = pptx.Presentation()
slide = deck.slides.add_slide(deck.slide_layouts[5])
slide.shapes.title.text = "Status"
slide.shapes.add_picture("chart.png", 0, 1500000)
deck.save("r.pptx")
import xlsxwriter
book = xlsxwriter.Workbook("r.xlsx")
sheet = book.add_worksheet("Overdue")
sheet.write_row(0, 0, ["Number", "Title"])
sheet.write_row(1, 0, [1, "Fix the gate"])
book.close()
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate
SimpleDocTemplate("r.pdf").build([Paragraph("Status", getSampleStyleSheet()["Normal"])])
import openpyxl
import pdfplumber
assert openpyxl.load_workbook("r.xlsx")["Overdue"]["B2"].value == "Fix the gate"
assert docx.Document("r.docx").paragraphs[1].text == "\\u20b9500 per user"
assert pptx.Presentation("r.pptx").slides[0].shapes.title.text == "Status"
with pdfplumber.open("r.pdf") as pdf:
    assert "Status" in pdf.pages[0].extract_text()
from PIL import Image
assert Image.open("chart.png").format == "PNG"
import networkx as nx
assert nx.dag_longest_path(nx.DiGraph([("a", "b"), ("b", "c")])) == ["a", "b", "c"]
from bs4 import BeautifulSoup
assert BeautifulSoup("<p>x</p>", "lxml").p.text == "x"
import jinja2
assert jinja2.Template("{{{{ n }}}}").render(n=1) == "1"
from scipy import stats
assert round(float(stats.norm.cdf(0)), 2) == 0.5
print("files:", " ".join(sorted(f for f in __import__("os").listdir(".") if "." in f)))
print("smoke: ok")
PY
"""


@pytest.mark.sandbox_docker
def test_the_image_imports_every_package_and_makes_each_file(
    coding_sandbox_image: str,  # noqa: F811 — the fixture, imported above
) -> None:
    """As uid 1000, read-only, with no network: every module, and five files."""
    script = _SMOKE_SCRIPT.format(modules=tuple(_SANDBOX_MODULES.values()))
    result = _run_in_image(coding_sandbox_image, 1000, script)
    assert result.returncode == 0, f"stdout:\n{result.stdout}\nstderr:\n{result.stderr}"
    lines = result.stdout.splitlines()
    assert "smoke: ok" in lines, result.stdout
    files = next(line for line in lines if line.startswith("files: ")).split()[1:]
    assert {"chart.png", "r.docx", "r.pptx", "r.xlsx", "r.pdf"} <= set(files), files
